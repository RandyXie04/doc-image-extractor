# -*- coding: utf-8 -*-
import os
import tempfile
import unittest
import docx
import pypandoc

from src.scripts.md_to_docx import apply_table_styles_to_docx
from src.utils.path_helper import get_template_path


class TestMdTableStyles(unittest.TestCase):
    """
    驗證自動整稿管線中表格製作樣式：
    - 表格樣式設為「表格格線」 (Table Grid)
    - 表格內文字段落無任何特殊樣式設定 (移除 Compact 等 pStyle，套用 Normal 內文)
    """

    def setUp(self):
        self.temp_dir = tempfile.mkdtemp(prefix="test_tbl_")

    def tearDown(self):
        import shutil
        shutil.rmtree(self.temp_dir, ignore_errors=True)

    def test_apply_table_styles_with_template(self):
        docx_path = os.path.join(self.temp_dir, "test_with_template.docx")
        template_path = str(get_template_path())

        md_content = """# 測試標題

這是一般本文段落。

| 欄位 A | 欄位 B |
| :--- | ---: |
| 內容 **粗體** | 內容 *斜體* |
| 第一列 | 第二列 |
"""
        pypandoc.convert_text(
            md_content,
            "docx",
            format="gfm",
            outputfile=docx_path,
            extra_args=[f"--reference-doc={template_path}"] if os.path.exists(template_path) else []
        )

        # 執行表格樣式後處理
        apply_table_styles_to_docx(docx_path)

        doc = docx.Document(docx_path)
        self.assertEqual(len(doc.tables), 1)
        tbl = doc.tables[0]

        # 驗證表格樣式為 Table Grid (表格格線)
        if os.path.exists(template_path):
            self.assertEqual(tbl.style.name, "Table Grid")
            self.assertEqual(tbl.style.style_id, "afe")
        else:
            self.assertIn("TableGrid", tbl._tbl.tblPr.xml)

        # 驗證文字段落無特殊樣式設定 (移除 Compact，為預設 Normal，且無殘留 pStyle 標籤)
        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    self.assertEqual(p.style.name, "Normal")
                    pPr = p._p.pPr
                    if pPr is not None:
                        # 確保無任何 pStyle (例如 Compact)
                        pStyle = pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle")
                        self.assertIsNone(pStyle)

    def test_apply_table_styles_without_template(self):
        docx_path = os.path.join(self.temp_dir, "test_no_template.docx")

        md_content = """
| 項目 | 數量 |
| --- | --- |
| 蘋果 | 10 |
"""
        pypandoc.convert_text(
            md_content,
            "docx",
            format="gfm",
            outputfile=docx_path
        )

        apply_table_styles_to_docx(docx_path)

        doc = docx.Document(docx_path)
        self.assertEqual(len(doc.tables), 1)
        tbl = doc.tables[0]
        self.assertIn("TableGrid", tbl._tbl.tblPr.xml)

        for row in tbl.rows:
            for cell in row.cells:
                for p in cell.paragraphs:
                    self.assertEqual(p.style.name, "Normal")
                    pPr = p._p.pPr
                    if pPr is not None:
                        pStyle = pPr.find("{http://schemas.openxmlformats.org/wordprocessingml/2006/main}pStyle")
                        self.assertIsNone(pStyle)


if __name__ == "__main__":
    unittest.main()
