# -*- coding: utf-8 -*-
"""
測試文件圖片提取與灰階轉換功能 (DOCX / PDF)
- 驗證 Word (.docx) 刪除圖片後不會被重複提取
- 驗證含 Alpha 透明通道圖片轉灰階時不會填充黑底
- 驗證 PDF 圖片提取及灰階轉換
"""

import os
import shutil
import tempfile
import unittest
import zipfile
from PIL import Image

import docx
import fitz

from src.scripts.extract_images import (
    extract_from_docx,
    extract_from_pdf,
    convert_to_grayscale_helper,
    process_document_images,
)


class TestExtractImages(unittest.TestCase):
    def setUp(self):
        self.test_dir = tempfile.mkdtemp(prefix="test_extract_img_")

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def test_docx_deleted_image_not_extracted(self):
        """測試 Word 中已被刪除的圖片不會被提取"""
        # 1. 建立兩張測試圖片
        img1_path = os.path.join(self.test_dir, "img1.png")
        img2_path = os.path.join(self.test_dir, "img2.png")
        Image.new("RGB", (20, 20), (255, 0, 0)).save(img1_path)
        Image.new("RGB", (20, 20), (0, 255, 0)).save(img2_path)

        # 2. 建立含有兩張圖的 Word 文件
        two_img_docx = os.path.join(self.test_dir, "two.docx")
        doc = docx.Document()
        doc.add_paragraph("Img1").add_run().add_picture(img1_path)
        doc.add_paragraph("Img2").add_run().add_picture(img2_path)
        doc.save(two_img_docx)

        # 3. 模擬使用者在 Word 刪除 Img2，Word ZIP 內部 media/image2.png 仍殘留，但 document.xml 移除
        deleted_docx = os.path.join(self.test_dir, "deleted.docx")
        with zipfile.ZipFile(two_img_docx, "r") as z_in:
            doc_xml = z_in.read("word/document.xml").decode("utf-8")
            p2_start = doc_xml.find("<w:p><w:r><w:t>Img2")
            p2_end = doc_xml.find("</w:p><w:sectPr") + len("</w:p>")
            doc_xml_modified = doc_xml[:p2_start] + doc_xml[p2_end:]

            with zipfile.ZipFile(deleted_docx, "w") as z_out:
                for item in z_in.infolist():
                    if item.filename == "word/document.xml":
                        z_out.writestr(item, doc_xml_modified.encode("utf-8"))
                    else:
                        z_out.writestr(item, z_in.read(item.filename))

        # 4. 驗證提取時只提取 active 的圖片 (只有 1 張，不含被刪除的 img2)
        out_dir = os.path.join(self.test_dir, "out_deleted")
        count = extract_from_docx(deleted_docx, out_dir)
        self.assertEqual(count, 1)
        files = os.listdir(out_dir)
        self.assertEqual(len(files), 1)
        self.assertEqual(files[0], "image1.png")

    def test_docx_transparent_png_grayscale_preserves_alpha(self):
        """測試含 Alpha 透明通道的 PNG 轉灰階時保留透明度，透明處不轉黑"""
        trans_png_path = os.path.join(self.test_dir, "trans.png")
        t_img = Image.new("RGBA", (20, 20), (0, 0, 0, 0))  # 全透明黑底
        t_img.putpixel((10, 10), (255, 0, 0, 255))          # 中央紅色不透明點
        t_img.save(trans_png_path)

        doc_path = os.path.join(self.test_dir, "trans_doc.docx")
        doc = docx.Document()
        doc.add_paragraph().add_run().add_picture(trans_png_path)
        doc.save(doc_path)

        out_dir = os.path.join(self.test_dir, "out_gray")
        count = extract_from_docx(doc_path, out_dir, to_grayscale=True)
        self.assertEqual(count, 1)

        result_path = os.path.join(out_dir, "image1.png")
        res = Image.open(result_path)

        # 模式應為 LA (灰階 + Alpha)
        self.assertEqual(res.mode, "LA")
        # 透明像素 (0,0) 的 Alpha 應為 0 (保持透明)，且亮度為 255 (白底防護，不填充黑)
        l_val, a_val = res.getpixel((0, 0))
        self.assertEqual(a_val, 0)
        self.assertEqual(l_val, 255)
        # 內容像素 (10,10) 應為不透明且轉為灰階
        l_dot, a_dot = res.getpixel((10, 10))
        self.assertEqual(a_dot, 255)
        self.assertEqual(l_dot, 76)  # 255*0.299 ≈ 76

    def test_pdf_transparent_png_grayscale(self):
        """測試 PDF 提取透明圖片轉灰階保留透明度"""
        trans_png_path = os.path.join(self.test_dir, "pdf_trans.png")
        t_img = Image.new("RGBA", (60, 60), (0, 0, 0, 0))
        t_img.putpixel((30, 30), (0, 255, 0, 255))
        t_img.save(trans_png_path)

        pdf_path = os.path.join(self.test_dir, "test.pdf")
        doc = fitz.open()
        page = doc.new_page()
        page.insert_image(fitz.Rect(50, 50, 150, 150), filename=trans_png_path)
        doc.save(pdf_path)
        doc.close()

        out_dir = os.path.join(self.test_dir, "out_pdf_gray")
        count = extract_from_pdf(pdf_path, out_dir, to_grayscale=True)
        self.assertEqual(count, 1)

        files = os.listdir(out_dir)
        self.assertEqual(len(files), 1)
        res = Image.open(os.path.join(out_dir, files[0]))
        self.assertEqual(res.mode, "LA")
        l_bg, a_bg = res.getpixel((0, 0))
        self.assertEqual(a_bg, 0)
        self.assertEqual(l_bg, 255)

    def test_grayscale_helper_unsupported_alpha_format(self):
        """測試轉為不支援透明度的格式 (如 JPG) 時，透明處合成白底而非黑底"""
        t_img = Image.new("RGBA", (10, 10), (0, 0, 0, 0))
        t_img.putpixel((5, 5), (0, 0, 255, 255))
        res = convert_to_grayscale_helper(t_img, ".jpg")
        self.assertEqual(res.mode, "L")
        # 透明處合成到白底後亮度應為 255
        self.assertEqual(res.getpixel((0, 0)), 255)
        # 藍色處轉灰階亮度應為 29 (255*0.114)
        self.assertEqual(res.getpixel((5, 5)), 29)

    def test_docx_duplicate_image_references_all_extracted(self):
        """測試 Word 中同一圖檔於不同位置被多次引用時，預設全部提取每一張"""
        img_path = os.path.join(self.test_dir, "shared.png")
        Image.new("RGB", (20, 20), (120, 180, 240)).save(img_path)

        doc_path = os.path.join(self.test_dir, "multi_ref.docx")
        doc = docx.Document()
        doc.add_paragraph("Paragraph 1").add_run().add_picture(img_path)
        doc.add_paragraph("Paragraph 2 without image")
        doc.add_paragraph("Paragraph 3").add_run().add_picture(img_path)
        doc.save(doc_path)

        # 1. 預設模式 (every_occurrence) 應提取 2 張圖
        out_every = os.path.join(self.test_dir, "out_every")
        count_every = extract_from_docx(doc_path, out_every)
        self.assertEqual(count_every, 2)
        files_every = sorted(os.listdir(out_every))
        self.assertEqual(files_every, ["image1.png", "image2.png"])

        # 2. 指定 unique 模式時僅提取 1 張圖
        out_unique = os.path.join(self.test_dir, "out_unique")
        count_unique = extract_from_docx(doc_path, out_unique, duplicate_policy="unique")
        self.assertEqual(count_unique, 1)
        files_unique = sorted(os.listdir(out_unique))
        self.assertEqual(files_unique, ["image1.png"])

    def test_docx_alternate_content_fallback_not_double_counted(self):
        """測試 Word 的 AlternateContent (Choice + Fallback) 不會被誤判為重複引用兩次"""
        img_path = os.path.join(self.test_dir, "single.png")
        Image.new("RGB", (20, 20), (200, 50, 50)).save(img_path)

        base_docx = os.path.join(self.test_dir, "base.docx")
        doc = docx.Document()
        doc.add_paragraph("Img").add_run().add_picture(img_path)
        doc.save(base_docx)

        # 將 drawing 封裝入 AlternateContent (含 Choice 與 Fallback)
        ac_docx = os.path.join(self.test_dir, "ac.docx")
        with zipfile.ZipFile(base_docx, "r") as z_in:
            doc_xml = z_in.read("word/document.xml").decode("utf-8")
            draw_start = doc_xml.find("<w:drawing>")
            draw_end = doc_xml.find("</w:drawing>") + len("</w:drawing>")
            drawing_xml = doc_xml[draw_start:draw_end]

            ac_xml = (
                f'<mc:AlternateContent xmlns:mc="http://schemas.openxmlformats.org/markup-compatibility/2006">'
                f'<mc:Choice Requires="wps">{drawing_xml}</mc:Choice>'
                f'<mc:Fallback>{drawing_xml}</mc:Fallback>'
                f'</mc:AlternateContent>'
            )
            doc_xml_modified = doc_xml[:draw_start] + ac_xml + doc_xml[draw_end:]

            with zipfile.ZipFile(ac_docx, "w") as z_out:
                for item in z_in.infolist():
                    if item.filename == "word/document.xml":
                        z_out.writestr(item, doc_xml_modified.encode("utf-8"))
                    else:
                        z_out.writestr(item, z_in.read(item.filename))

        out_ac = os.path.join(self.test_dir, "out_ac")
        count = extract_from_docx(ac_docx, out_ac)
        # 即使含 Choice 與 Fallback，單一圖片應仍只提取 1 張，不可誤計為 2 張
        self.assertEqual(count, 1)
        self.assertEqual(os.listdir(out_ac), ["image1.png"])

    def test_docx_process_document_images_pipeline(self):
        """測試 process_document_images 整合流程自動打包多引用圖片"""
        img_path = os.path.join(self.test_dir, "pipeline_img.png")
        Image.new("RGB", (20, 20), (0, 100, 200)).save(img_path)

        doc_path = os.path.join(self.test_dir, "pipeline_doc.docx")
        doc = docx.Document()
        doc.add_paragraph("P1").add_run().add_picture(img_path)
        doc.add_paragraph("P2").add_run().add_picture(img_path)
        doc.save(doc_path)

        zip_out = os.path.join(self.test_dir, "result.zip")
        temp_extract = os.path.join(self.test_dir, "temp_extract")
        count, out_zip = process_document_images(doc_path, zip_out, temp_extract)
        self.assertEqual(count, 2)
        self.assertTrue(os.path.exists(out_zip))

        with zipfile.ZipFile(out_zip, "r") as z:
            extracted_names = sorted(z.namelist())
            self.assertEqual(extracted_names, ["image1.png", "image2.png"])


if __name__ == "__main__":
    unittest.main()
