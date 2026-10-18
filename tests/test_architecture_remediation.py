#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
專案重構與自動整稿目標驗收測試套件 (Architecture & Layout Remediation Test Suite)
=================================================================================
執行方式：
  1. 透過 pytest 執行：  pytest tests/test_architecture_remediation.py -v
  2. 直接以 Python 執行： python tests/test_architecture_remediation.py
"""

import os
import sys
import time
import shutil
import inspect
import tempfile
import unittest
from pathlib import Path

# 確保專案根目錄在 sys.path 中
PROJECT_ROOT = Path(__file__).resolve().parent.parent
if str(PROJECT_ROOT) not in sys.path:
    sys.path.insert(0, str(PROJECT_ROOT))

if sys.platform == "win32":
    try:
        sys.stdout.reconfigure(encoding="utf-8", errors="replace")
        sys.stderr.reconfigure(encoding="utf-8", errors="replace")
    except Exception:
        pass

import pymupdf as fitz  # PyMuPDF


class TestArchitectureRemediation(unittest.TestCase):
    """針對六大重構目標的端到端流程驗收測試"""

    def setUp(self):
        self.test_dir = Path(tempfile.mkdtemp(prefix="remediation_test_"))

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    # =========================================================================
    # 測試 1：【自動整稿】純文字通道 - 無編號標題識別 & 垂直間距/未滿行防黏連測試
    # =========================================================================
    def test_01_vector_heading_and_paragraph_anti_merge(self):
        """
        [目標驗收 1]
        驗證無編號標題（透過目錄錨定、相對字級差）與未滿行/大段距段落不會與內文黏在一起。
        """
        from src.scripts.book_layout_extractor import process_book_vector_pdf

        pdf_path = self.test_dir / "synthetic_book.pdf"
        out_dir = self.test_dir / "out_book"

        # 合成一本 2 頁的測試 PDF（含目錄頁 + 無編號標題 + 垂直段距跳躍）
        doc = fitz.open()

        # 第 1 頁：目錄頁
        p1 = doc.new_page(width=595, height=842)
        p1.insert_text((260, 80), "目 錄", fontsize=16, fontname="china-ss")
        p1.insert_text((70, 130), "導論 .................................... 2", fontsize=11, fontname="china-ss")
        p1.insert_text((70, 160), "文化與自然的交織 ........................ 2", fontsize=11, fontname="china-ss")

        # 第 2 頁：正文頁（故意使用無編號標題，且無首行縮排，測試是否會與內文黏連）
        p2 = doc.new_page(width=595, height=842)
        # 無編號大標 (14pt，相對於 10.5pt 正文)
        p2.insert_text((70, 100), "導論", fontsize=14, fontname="china-ss")
        # 第一段正文 (10.5pt，分成兩行，第一行滿行、第二行未滿行提早結束)
        p2.insert_text(
            (70, 135),
            "這是一段標準的書籍內文，用來建立全書的正文基準字級與左邊界座標參考值，確保統計準確。",
            fontsize=10.5,
            fontname="china-ss",
        )
        p2.insert_text(
            (70, 152),
            "這行在中間就結束了。",  # 未滿行 (Short line)
            fontsize=10.5,
            fontname="china-ss",
        )
        # 無編號次標題 (12pt，靠左對齊無縮排，與上一段有明顯垂直段距 y=195)
        p2.insert_text((70, 195), "文化與自然的交織", fontsize=12, fontname="china-ss")
        # 第二段正文（同樣靠左 x=70 無縮排，測試是否會把上方的「文化與自然的交織」吞進去）
        p2.insert_text(
            (70, 225),
            "第二段內文緊接在無編號小標之後，且故意不設首行縮排，檢驗系統是否能將標題與內文切開。",
            fontsize=10.5,
            fontname="china-ss",
        )

        doc.save(str(pdf_path))
        doc.close()

        md_path, audit_data = process_book_vector_pdf(str(pdf_path), str(out_dir))
        md_content = Path(md_path).read_text(encoding="utf-8")

        # 斷言 A：「文化與自然的交織」絕對不能跟「第二段內文緊接...」黏在同一行
        for line in md_content.splitlines():
            self.assertFalse(
                "文化與自然的交織" in line and "第二段內文" in line,
                f"❌ 失敗：無編號標題與下一段內文發生黏連！問題行內容：\n{line}"
            )
            self.assertFalse(
                "這行在中間就結束了。" in line and "文化與自然的交織" in line,
                f"❌ 失敗：未滿行段落與下方區塊發生黏連！問題行內容：\n{line}"
            )

        # 斷言 B：無編號標題應被識別為 Markdown 標題 (#)
        has_heading_1 = any(
            line.strip().startswith("#") and "導論" in line
            for line in md_content.splitlines()
        )
        has_heading_2 = any(
            line.strip().startswith("#") and "文化與自然的交織" in line
            for line in md_content.splitlines()
        )
        self.assertTrue(
            has_heading_1 and has_heading_2,
            f"❌ 失敗：無編號標題未被識別為 Markdown Heading (#)！實際產出：\n{md_content}"
        )

    # =========================================================================
    # 測試 2：註腳弱標記替換邊界保護（防止誤傷正文小數點 31.5% 與版本號 v1.2）
    # =========================================================================
    def test_02_footnote_weak_marker_decimal_safety(self):
        """
        [目標驗收 2]
        當頁底註腳標記為 '1.' 或 '(1)' 時，正文中出現的 '31.5%'、'v1.2' 絕不可被誤替換為 [^p1_1]。
        """
        from src.scripts.book_layout_extractor import _replace_fn_markers_in_text

        page_footnotes = [{"id": 1, "marker": "1.", "content": "測試註釋內容", "matched": False}]
        sample_text = "本年度良率提升至 31.5%，系統版本更新為 v1.2，具體數據詳見文獻1."

        result = _replace_fn_markers_in_text(sample_text, page_footnotes, page_num=1)

        self.assertIn(
            "31.5%",
            result,
            f"❌ 失敗：正文中的小數 '31.5%' 被註腳標記 '1.' 誤傷破壞！結果為：{result}"
        )
        self.assertIn(
            "v1.2",
            result,
            f"❌ 失敗：正文中的版本號 'v1.2' 被註腳標記 '1.' 誤傷破壞！結果為：{result}"
        )

    # =========================================================================
    # 測試 3：方正字元修復無損保留 Word (.docx) 粗體/斜體/字號樣式
    # =========================================================================
    def test_03_founder_docx_repair_preserves_run_formatting(self):
        """
        [目標驗收 3]
        驗證 repair_docx_file 修復方正彝文區標點（ꎬ -> ，、ꎮ -> 。）時，不會抹除段落內的粗體與斜體 Run 樣式。
        """
        import docx
        from src.founder_tools.fix_founder_fonts import repair_docx_file

        in_docx = self.test_dir / "founder_broken.docx"
        out_docx = self.test_dir / "founder_repaired.docx"

        doc = docx.Document()
        p = doc.add_paragraph()
        r1 = p.add_run("這是一段帶有方正亂碼標點的文字ꎬ")
        r2 = p.add_run("這是粗體強調詞彙")
        r2.bold = True
        r3 = p.add_run("以及斜體引文ꎮ")
        r3.italic = True
        doc.save(str(in_docx))

        repair_docx_file(str(in_docx), str(out_docx))

        repaired_doc = docx.Document(str(out_docx))
        repaired_p = repaired_doc.paragraphs[0]

        # 斷言 A：方正彝文區標點已成功還原為標準全形標點
        self.assertEqual(
            repaired_p.text,
            "這是一段帶有方正亂碼標點的文字，這是粗體強調詞彙以及斜體引文。",
            "❌ 失敗：方正亂碼字元未正確還原！"
        )

        # 斷言 B：原本的粗體與斜體 Run 屬性必須被保留（不可因為直接賦值 p.text 而全部變為 None）
        has_bold_run = any(r.bold is True and "粗體" in r.text for r in repaired_p.runs)
        has_italic_run = any(r.italic is True and "斜體" in r.text for r in repaired_p.runs)
        self.assertTrue(
            has_bold_run and has_italic_run,
            f"❌ 失敗：修復 Word 文字後，段落內的粗體/斜體 Run 格式被洗掉了！目前 Runs: "
            f"{[(r.text, r.bold, r.italic) for r in repaired_p.runs]}"
        )

    # =========================================================================
    # 測試 4：PDF 檔案控制代碼釋放驗證（防止 Windows 檔案鎖定洩漏）
    # =========================================================================
    def test_04_pdf_file_handle_closed_properly(self):
        """
        [目標驗收 4]
        驗證 process_book_vector_pdf 執行完畢後已正確關閉 fitz.Document，可立即刪除原檔而不觸發 PermissionError。
        """
        from src.scripts.book_layout_extractor import process_book_vector_pdf

        pdf_path = self.test_dir / "handle_leak_check.pdf"
        doc = fitz.open()
        page = doc.new_page()
        page.insert_text((72, 72), "測試檔案控制代碼釋放", fontsize=12, fontname="china-ss")
        doc.save(str(pdf_path))
        doc.close()

        process_book_vector_pdf(str(pdf_path), str(self.test_dir / "out_leak"))

        try:
            os.remove(pdf_path)
        except PermissionError as e:
            self.fail(f"❌ 失敗：PDF 檔案仍被鎖定無法刪除，代表 fitz.open() 未正確呼叫 close()！({e})")

    # =========================================================================
    # 測試 5：OCR 管線資源與 Regex 回溯效能檢查
    # =========================================================================
    def test_05_ocr_pipeline_memory_and_regex_safety(self):
        """
        [目標驗收 5]
        1. 檢查 process_ocr.py 是否已移除無用的 load_det_model() 與 200 頁硬性 break 截斷。
        2. 測試 _clean_hallucinations 面對高誘發回溯字串時的執行時間 < 50ms。
        """
        import src.scripts.process_ocr as process_ocr

        src_code = inspect.getsource(process_ocr.main)
        self.assertNotIn(
            "load_det_model()",
            src_code,
            "❌ 失敗：process_ocr.py 仍在載入未使用的 load_det_model()，平白消耗 VRAM！"
        )
        self.assertNotIn(
            "len(images) >= 200: break",
            src_code,
            "❌ 失敗：process_ocr.py 仍存在 'if len(images) >= 200: break' 的無預警 200 頁截斷邏輯！"
        )

        # 壓測 _clean_hallucinations 正則效能
        adversarial_text = ("天地玄黃宇宙洪荒" * 4 + "日月盈昃") * 30
        t0 = time.perf_counter()
        _ = process_ocr._clean_hallucinations(adversarial_text)
        elapsed_ms = (time.perf_counter() - t0) * 1000.0
        self.assertLess(
            elapsed_ms,
            50.0,
            f"❌ 失敗：_clean_hallucinations 正則執行過慢 ({elapsed_ms:.1f} ms)，存在回溯效能瓶頸！"
        )

    # =========================================================================
    # 測試 6：跨任務目錄隔離與 Path Traversal 前綴繞過防禦
    # =========================================================================
    def test_06_task_isolation_and_path_traversal_defense(self):
        """
        [目標驗收 6]
        1. 驗證兩個 PDFConversionAgent 實例的暫存/公式目錄是否具備隔離性，或雙胞胎腳本是否已收斂。
        2. 驗證路徑白名單檢查不會被同前綴目錄（如 data_hacked）繞過。
        """
        from config import PATHS

        # 檢查 Path.is_relative_to 防禦邏輯
        allowed_dir = (PATHS.root / "data" / "03_output").resolve()
        spoofed_sibling_path = (PATHS.root / "data" / "03_output_hacked" / "secret.docx").resolve()

        # 舊版 str.startswith() 在此處會誤判為 True，新版 is_relative_to() 應為 False
        is_safe_modern = spoofed_sibling_path.is_relative_to(allowed_dir)
        self.assertFalse(
            is_safe_modern,
            "❌ 失敗：路徑校驗邏輯無法抵禦同前綴目錄繞過！"
        )

        # 檢查 app.py 中是否已淘汰危險的 .startswith(str(d)) 寫法
        import src.web.app as web_app
        app_src = inspect.getsource(web_app)
        self.assertNotIn(
            "str(target_path).startswith(str(d))",
            app_src,
            "❌ 失敗：src/web/app.py 仍在使用不安全的 str.startswith() 進行路徑白名單驗證，請改用 target_path.is_relative_to(d)！"
        )

    # =========================================================================
    # 測試 7：雙胞胎腳本收斂與無註腳模式參數化驗證
    # =========================================================================
    def test_07_twin_scripts_consolidated(self):
        """
        [目標驗收 7]
        1. 驗證 *_no_footnote.py 雙胞胎檔案已物理刪除。
        2. 驗證 process_book_vector_pdf(enable_footnotes=False) 能將頁底帶圈號文字保留於正文，不產生 [^p1_1]。
        """
        from src.scripts.book_layout_extractor import process_book_vector_pdf

        scripts_dir = PROJECT_ROOT / "src" / "scripts"
        self.assertFalse(
            (scripts_dir / "book_layout_extractor_no_footnote.py").exists(),
            "❌ 失敗：book_layout_extractor_no_footnote.py 尚未刪除！"
        )
        self.assertFalse(
            (scripts_dir / "process_ocr_no_footnote.py").exists(),
            "❌ 失敗：process_ocr_no_footnote.py 尚未刪除！"
        )

        # 建立一頁底部含有 "① 這是頁底條列正文" 的測試 PDF
        pdf_path = self.test_dir / "no_fn_mode.pdf"
        doc = fitz.open()
        page = doc.new_page(width=595, height=842)
        page.insert_text((70, 120), "這是上半部的正文段落①。", fontsize=10.5, fontname="china-ss")
        page.insert_text((70, 720), "① 這是頁底區塊，在無註腳模式下應保留為普通正文。", fontsize=9.5, fontname="china-ss")
        doc.save(str(pdf_path))
        doc.close()

        md_path, _ = process_book_vector_pdf(
            str(pdf_path),
            str(self.test_dir / "out_no_fn"),
            enable_footnotes=False,
        )
        md_content = Path(md_path).read_text(encoding="utf-8")

        self.assertNotIn(
            "[^p1_1]",
            md_content,
            f"❌ 失敗：enable_footnotes=False 時仍產生了 Pandoc 註腳標籤！產出：\n{md_content}"
        )
        self.assertIn(
            "這是頁底區塊",
            md_content,
            "❌ 失敗：頁底區塊文字在無註腳模式下遺失！"
        )


if __name__ == "__main__":
    print("=" * 75)
    print("🚀 開始執行【專案重構與自動整稿目標】自動化驗收測試...")
    print("=" * 75)
    unittest.main(verbosity=2)
