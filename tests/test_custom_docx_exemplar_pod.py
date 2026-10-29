# -*- coding: utf-8 -*-
"""
Tests for POD Studio handling non-default custom DOCX exemplars and page ranges.
"""
import pytest
import tempfile
import shutil
from pathlib import Path
import docx

from src.pod_studio.exemplar_learner import ExemplarLearner, DiffAligner
from src.pod_studio.blueprint_schema import BookBlueprint, StyleTriggerRule
from src.pod_studio.section_ast import SectionUnit, ParagraphUnit
from src.pod_studio.propagation_engine import PropagationEngine
from src.pod_studio.pod_validator import PodValidator
from src.pod_studio.pod_service import (
    extract_raw_from_pdf_ast,
    compile_blueprint_from_seed,
    propagate_and_validate_session,
)


class TestCustomDocxExemplarPOD:
    """Verifies that user-provided non-default DOCX styles are learned and applied accurately."""

    def test_structural_pattern_induction_from_custom_docx(self, tmp_path):
        """Ensures non-default Word styles like '一、', '（）', '方劑【組成】' are induced."""
        doc = docx.Document()
        # Add custom style mock paragraphs
        p1 = doc.add_paragraph("一、發熱的定義")
        p1.style = doc.styles.add_style("一、", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)

        p2 = doc.add_paragraph("二、發熱的機制")
        p2.style = "一、"

        p3 = doc.add_paragraph("（一）體溫調節")
        p3.style = doc.styles.add_style("（）", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)

        p4 = doc.add_paragraph("（二）散熱障礙")
        p4.style = "（）"

        p5 = doc.add_paragraph("【組成】柴胡12g 黃芩9g")
        p5.style = doc.styles.add_style("方劑【組成】", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)

        p6 = doc.add_paragraph("【組成】半夏15g 陳皮15g")
        p6.style = "方劑【組成】"

        gold_path = tmp_path / "custom_seed.docx"
        doc.save(str(gold_path))

        # Raw document
        raw_doc = docx.Document()
        raw_doc.add_paragraph("## 第一章 概述")
        raw_doc.add_paragraph("一、發熱的定義")
        raw_doc.add_paragraph("正文描述內容。")
        raw_path = tmp_path / "raw.docx"
        raw_doc.save(str(raw_path))

        bp = ExemplarLearner.learn_from_docx_pair(raw_path, gold_path)

        styles_learned = set(bp.template_styles_used)
        assert "一、" in styles_learned
        assert "（）" in styles_learned
        assert "方劑【組成】" in styles_learned

        # Verify induced rules
        rule_targets = {r.target_docx_style for r in bp.style_rules if r.action == "apply_style"}
        assert "一、" in rule_targets
        assert "（）" in rule_targets
        assert "方劑【組成】" in rule_targets

    def test_propagation_applies_to_preamble_and_preserves_cjk(self):
        """Verifies that preamble paragraphs receive custom style transformations."""
        blueprint = BookBlueprint(
            book_name="test_book",
            book_type="general",
            template_styles_used=["Normal", "一、", "方劑【組成】"],
            style_rules=[
                StyleTriggerRule(
                    rule_id="r1",
                    action="apply_style",
                    target_docx_style="一、",
                    trigger_regex=r"^[一二三四五六七八九十百]+[、.· ]",
                    priority=15,
                ),
                StyleTriggerRule(
                    rule_id="r2",
                    action="apply_style",
                    target_docx_style="方劑【組成】",
                    trigger_regex=r"^【(?:组成|組成)】",
                    priority=10,
                ),
            ],
        )

        engine = PropagationEngine(blueprint)
        paras = [
            ParagraphUnit(para_id="p001", raw_text="普通開場白文字", page_num=1),
            ParagraphUnit(para_id="p002", raw_text="六、骨蒸熱", page_num=1),
            ParagraphUnit(para_id="p003", raw_text="【組成】半夏15g 陳皮15g", page_num=1),
        ]
        preamble = SectionUnit(
            section_id="sec_000",
            title="__PREAMBLE__",
            level=0,
            paragraphs=paras,
            raw_content="\n\n".join(p.raw_text for p in paras),
        )

        processed = engine.process_section(preamble)

        assert processed.paragraphs[0].custom_style in ("", "Normal")
        assert processed.paragraphs[1].custom_style == "一、"
        assert processed.paragraphs[2].custom_style == "方劑【組成】"

        # Verify CJK conservation
        validator = PodValidator(blueprint=blueprint)
        violations, cjk_res = validator.validate_section(processed)
        assert len(violations) == 0
        assert cjk_res.status == "GREEN"

    def test_alien_fonts_json_style_sanitized(self):
        """Verifies that custom_style tags not defined in template.docx are sanitized."""
        blueprint = BookBlueprint(
            book_name="clean_book",
            book_type="general",
            template_styles_used=["Normal", "一、"],
            style_rules=[],
        )

        engine = PropagationEngine(blueprint)
        sec = SectionUnit(
            section_id="sec_001",
            title="測試篇",
            level=2,
            paragraphs=[
                ParagraphUnit(para_id="p001", raw_text="普通段落", custom_style="標粗", page_num=1),
            ],
        )

        processed = engine.process_section(sec)
        # "標粗" is not in template_styles_used, so it must be sanitized to empty
        assert processed.paragraphs[0].custom_style == ""
