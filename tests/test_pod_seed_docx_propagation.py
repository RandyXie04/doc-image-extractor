# -*- coding: utf-8 -*-
"""
Tests for automatic completion of subsequent unmodified paragraphs in seed docx (Phase 3).
Verifies that:
1. ExemplarLearner learns generalized style rules, keyword triggers, and boundary from seed DOCX.
2. PropagationEngine applies hybrid deterministic + AI rules to subsequent paragraphs in the same DOCX.
3. The resulting Book_POD_Final.docx preserves human pre-edits and automatically styles subsequent paragraphs.
"""
import shutil
import tempfile
from pathlib import Path
import docx
import pytest

from src.pod_studio.exemplar_learner import ExemplarLearner
from src.pod_studio.propagation_engine import PropagationEngine
from src.pod_studio.section_ast import SectionAST, SectionUnit, ParagraphUnit
from src.pod_studio.pod_service import compile_blueprint_from_seed, propagate_and_validate_session


class TestSeedDocxSubsequentPropagation:
    def test_learn_from_real_exemplar_docx(self):
        """Verifies rule induction on 常见肿瘤中医诊治心悟_人工預修前.docx."""
        seed_path = Path("data/01_input/常见肿瘤中医诊治心悟_人工預修前.docx")
        if not seed_path.exists():
            pytest.skip("data/01_input/常见肿瘤中医诊治心悟_人工預修前.docx not found")

        bp = ExemplarLearner.learn_from_docx_pair(raw_docx_path="", golden_docx_path=seed_path)

        # Verify styles learned
        assert "藥方特點" in bp.template_styles_used
        assert "方劑組成1" in bp.template_styles_used
        assert "方劑組成2" in bp.template_styles_used
        assert "楷體2" in bp.template_styles_used
        assert "Heading 3" in bp.template_styles_used

        # Verify metadata
        assert bp.metadata.get("seed_last_modified_idx", 0) >= 150
        assert len(bp.metadata.get("style_definitions", {})) >= 5

        # Verify inferred rules
        rule_targets = {r.target_docx_style for r in bp.style_rules if r.action == "apply_style"}
        assert "藥方特點" in rule_targets
        assert "方劑組成1" in rule_targets
        assert "方劑組成2" in rule_targets
        assert "楷體2" in rule_targets

    def test_end_to_end_seed_docx_completion(self, tmp_path):
        """Verifies that Stage 3 automatically refines subsequent paragraphs in the same docx."""
        session_dir = tmp_path / "test_session"
        session_dir.mkdir(parents=True, exist_ok=True)

        seed_doc = docx.Document()
        seed_doc.styles.add_style("藥方特點", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
        seed_doc.styles.add_style("方劑組成1", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
        seed_doc.styles.add_style("方劑組成2", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)
        seed_doc.styles.add_style("楷體2", docx.enum.style.WD_STYLE_TYPE.PARAGRAPH)

        # Human pre-edited section (paragraphs 0..4)
        p0 = seed_doc.add_paragraph("方剂出处　《伤寒论》")
        p0.style = "藥方特點"
        p1 = seed_doc.add_paragraph("原文记载")
        p1.style = "藥方特點"
        p2 = seed_doc.add_paragraph("《伤寒论》：“太阳病，发汗后...”")
        p2.style = "楷體2"
        p3 = seed_doc.add_paragraph("方剂组成")
        p3.style = "方劑組成1"
        p4 = seed_doc.add_paragraph("黄芪30g，柴胡12g，白术15g。")
        p4.style = "方劑組成2"

        # Subsequent un-modified section (paragraphs 5..10, all Normal in seed doc)
        p5 = seed_doc.add_paragraph("方剂出处　《金匮要略》")
        p6 = seed_doc.add_paragraph("原文记载")
        p7 = seed_doc.add_paragraph("《金匮要略》：“血痹虚劳...”")
        p8 = seed_doc.add_paragraph("方剂组成")
        p9 = seed_doc.add_paragraph("熟地黄30g，山药30g，山茱萸15g。")
        p10 = seed_doc.add_paragraph("普通患者内文描述说明。")

        seed_file = session_dir / "Book_seed.docx"
        seed_doc.save(str(seed_file))

        # Also create Book_raw_ast.json for AST validation
        sections = [
            SectionUnit(
                section_id="sec_001",
                title="第一章 測試章節",
                level=1,
                paragraphs=[
                    ParagraphUnit(para_id=f"sec_001_p{i:03d}", raw_text=p.text, page_num=1)
                    for i, p in enumerate(seed_doc.paragraphs)
                ],
                raw_content="\n\n".join(p.text for p in seed_doc.paragraphs),
            )
        ]
        SectionAST.dump_book_ast_json(sections, session_dir / "Book_raw_ast.json")

        # Compile blueprint from seed
        blueprint = compile_blueprint_from_seed(session_dir, seed_file)
        assert blueprint.metadata.get("seed_last_modified_idx") == 4

        # Execute Stage 3 propagate_and_validate_session with a fast mock AI client for any unclassified
        mock_ai = lambda prompt: '[{"id": 10, "style": "Normal"}]'

        final_docx, report, bp_out = propagate_and_validate_session(
            session_dir=session_dir,
            custom_blueprint=blueprint,
            ai_client=mock_ai,
        )

        assert final_docx.exists()
        out_doc = docx.Document(str(final_docx))
        assert len(out_doc.paragraphs) == 11

        # Human pre-edited paragraphs must remain unchanged
        assert out_doc.paragraphs[0].style.name == "藥方特點"
        assert out_doc.paragraphs[1].style.name == "藥方特點"
        assert out_doc.paragraphs[2].style.name == "楷體2"
        assert out_doc.paragraphs[3].style.name == "方劑組成1"
        assert out_doc.paragraphs[4].style.name == "方劑組成2"

        # Subsequent paragraphs must be automatically converted according to learned rules!
        assert out_doc.paragraphs[5].style.name == "藥方特點"  # "方剂出处　《金匮要略》"
        assert out_doc.paragraphs[6].style.name == "藥方特點"  # "原文记载"
        assert out_doc.paragraphs[7].style.name == "楷體2"    # "《金匮要略》：“血痹虚劳...”"
        assert out_doc.paragraphs[8].style.name == "方劑組成1"  # "方剂组成"
        assert out_doc.paragraphs[9].style.name == "方劑組成2"  # "熟地黄30g，山药30g..."
        assert out_doc.paragraphs[10].style.name in ("Normal", "Normal Table")  # Regular text
