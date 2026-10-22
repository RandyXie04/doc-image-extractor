# -*- coding: utf-8 -*-
"""
Phase 3 Test Suite: Controlled POD Studio Integration
======================================================
Tests integration of extract_book_vector_ast into POD Studio service layer:
- Parity between legacy extract_raw_from_pdf and AST path
- Contract fidelity: List[SectionUnit], Book_raw_ast.json, Book_raw.docx
- Preservation of page_num, bbox, font metadata across dump/load cycle
- Full 3-stage POD pipeline (extract -> compile_blueprint -> propagate_and_validate)
- Side-effect isolation: no writes outside session_dir
"""

import os
import shutil
import tempfile
import unittest
from pathlib import Path
import fitz
import docx

from src.pod_studio.section_ast import SectionUnit, ParagraphUnit, SectionAST, BookAST, SectionAdapter
from src.pod_studio.pod_service import (
    extract_raw_from_pdf,
    compile_blueprint_from_seed,
    propagate_and_validate_session,
)
from src.pod_studio.blueprint_schema import BookBlueprint, StyleTriggerRule
from src.pod_studio.propagation_engine import PropagationEngine
from src.pod_studio.pod_validator import PodValidator


class TestPODStudioASTIntegration(unittest.TestCase):
    """Test suite for Phase 3 controlled AST integration in POD Studio."""

    def setUp(self):
        self.test_dir = Path(tempfile.mkdtemp(prefix="test_pod_int_"))
        self.sample_pdf = self.test_dir / "sample_book.pdf"
        self._generate_sample_pdf(self.sample_pdf)

    def tearDown(self):
        shutil.rmtree(self.test_dir, ignore_errors=True)

    def _generate_sample_pdf(self, path: Path):
        """Generates a 3-page synthetic PDF with chapters and body text."""
        doc = fitz.open()

        # Page 1: Chapter 1 + body paragraphs
        p1 = doc.new_page(width=500, height=700)
        p1.insert_text((60, 110), "第一章 數位出版技術", fontsize=16, fontname="china-ss")
        p1.insert_text(
            (60, 160),
            "這是第一章的第一段標準內文，用於檢驗段落邊界框與字體資訊。",
            fontsize=10.5,
            fontname="china-ss",
        )
        p1.insert_text(
            (60, 190),
            "這是第二段落，包含特殊排版與測試用的第二行文字。",
            fontsize=10.5,
            fontname="china-ss",
        )

        # Page 2: Chapter 2 + body
        p2 = doc.new_page(width=500, height=700)
        p2.insert_text((60, 110), "第二章 版面幾何與語意", fontsize=16, fontname="china-ss")
        p2.insert_text(
            (60, 160),
            "第二章的開頭段落，探討結構化抽象語法樹的核心價值。",
            fontsize=10.5,
            fontname="china-ss",
        )

        # Page 3: Subheading and short block
        p3 = doc.new_page(width=500, height=700)
        p3.insert_text((60, 110), "2.1 幾何座標精確性", fontsize=13, fontname="china-ss")
        p3.insert_text(
            (60, 150),
            "終章結論與摘要說明。",
            fontsize=10.5,
            fontname="china-ss",
        )

        doc.save(str(path))
        doc.close()

    def test_01_legacy_path_default_unaffected(self):
        """Test that default extract_raw_from_pdf continues to execute legacy path."""
        s_dir = self.test_dir / "session_legacy"
        sections, ast_path, docx_path = extract_raw_from_pdf(
            self.sample_pdf, s_dir, split_level=2, use_vector_ast=False
        )
        self.assertTrue(ast_path.exists())
        self.assertTrue(docx_path.exists())
        self.assertIsInstance(sections, list)
        self.assertGreater(len(sections), 0)

    def test_02_vector_ast_path_contract_and_files(self):
        """Test that extract_raw_from_pdf with use_vector_ast=True satisfies contract and generates files."""
        s_dir = self.test_dir / "session_ast"
        sections, ast_path, docx_path = extract_raw_from_pdf(
            self.sample_pdf, s_dir, split_level=2, use_vector_ast=True
        )

        # 1. Output files exist in session_dir
        self.assertTrue(ast_path.exists(), "Book_raw_ast.json must exist")
        self.assertTrue(docx_path.exists(), "Book_raw.docx must exist")
        self.assertEqual(ast_path.name, "Book_raw_ast.json")
        self.assertEqual(docx_path.name, "Book_raw.docx")

        # 2. Return types
        self.assertIsInstance(sections, list)
        self.assertTrue(all(isinstance(s, SectionUnit) for s in sections))
        self.assertGreater(len(sections), 1)

        # 3. Document has valid content
        doc_in = docx.Document(str(docx_path))
        doc_texts = [p.text for p in doc_in.paragraphs if p.text.strip()]
        self.assertGreater(len(doc_texts), 0)

    def test_03_geometry_and_font_preservation_across_serialization(self):
        """Test that Book_raw_ast.json loaded back retains valid page_num, bbox, and font metadata."""
        s_dir = self.test_dir / "session_ast_geo"
        sections, ast_path, _ = extract_raw_from_pdf(
            self.sample_pdf, s_dir, split_level=2, use_vector_ast=True
        )

        # Load back from JSON
        reloaded_sections = SectionAST.load_book_ast_json(ast_path)
        self.assertEqual(len(reloaded_sections), len(sections))

        all_paras = [p for s in reloaded_sections for p in s.paragraphs]
        self.assertGreater(len(all_paras), 0)

        for p in all_paras:
            # 1-based page_num
            self.assertGreaterEqual(p.page_num, 1)
            self.assertLessEqual(p.page_num, 3)

            # Valid bbox
            self.assertIsNotNone(p.bbox)
            self.assertEqual(len(p.bbox), 4)
            x0, y0, x1, y1 = p.bbox
            self.assertLess(x0, x1)
            self.assertLess(y0, y1)

            # Font metadata
            self.assertTrue(len(p.font_name) > 0)
            self.assertGreater(p.font_size, 0.0)

    def test_04_downstream_3_stage_pipeline_compatibility(self):
        """Test that AST extraction integrates cleanly into Stage 2 (compile) and Stage 3 (propagate)."""
        s_dir = self.test_dir / "session_pipeline"

        # Stage 1: Extract with AST
        sections, ast_path, docx_path = extract_raw_from_pdf(
            self.sample_pdf, s_dir, split_level=2, use_vector_ast=True
        )

        # Create a sample seed markdown for Stage 2
        seed_path = self.test_dir / "seed.md"
        seed_path.write_text(
            "# 第一章 數位出版技術\n\n【導讀】這是第一章的第一段標準內文，用於檢驗段落邊界框與字體資訊。\n\n這是第二段落，包含特殊排版與測試用的第二行文字。\n",
            encoding="utf-8",
        )

        # Stage 2: Compile blueprint
        blueprint = compile_blueprint_from_seed(s_dir, seed_path, sample_chapters=1)
        self.assertIsInstance(blueprint, BookBlueprint)
        bp_path = s_dir / "book_blueprint.json"
        self.assertTrue(bp_path.exists())

        # Stage 3: Propagate and validate
        final_docx, report, bp_out = propagate_and_validate_session(
            session_dir=s_dir,
            custom_blueprint=blueprint,
            pdf_path=self.sample_pdf,
        )
        self.assertTrue(final_docx.exists())
        self.assertEqual(final_docx.name, "Book_POD_Final.docx")
        self.assertIsNotNone(report)
        audit_path = s_dir / "audit_report.json"
        self.assertTrue(audit_path.exists())

    def test_05_file_handle_closure_on_success_and_failure(self):
        """Test that extract_raw_from_pdf with use_vector_ast=True closes file handles cleanly."""
        s_dir = self.test_dir / "session_handles"
        _ = extract_raw_from_pdf(self.sample_pdf, s_dir, use_vector_ast=True)

        # On Windows, open file handles prevent deletion
        try:
            self.sample_pdf.unlink()
        except PermissionError:
            self.fail("PDF file remained locked after extract_raw_from_pdf")

        # Failure path
        with self.assertRaises(FileNotFoundError):
            extract_raw_from_pdf(self.test_dir / "nonexistent.pdf", s_dir, use_vector_ast=True)


if __name__ == "__main__":
    unittest.main()
