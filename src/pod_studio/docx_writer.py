# -*- coding: utf-8 -*-
"""
DocxSerializer (Deep Module for OpenXML Document Assembly)
==========================================================
Encapsulates all OpenXML / python-docx assembly complexity:
- Template styles preservation (clearing body paragraphs while retaining styles.xml)
- Heading hierarchy mapping
- Markdown table parsing and formatting
- Image embedding with scale adjustment
- Pandoc custom-style fence (::: {custom-style="..."}) translation
"""

from __future__ import annotations

import os
import re
from pathlib import Path
from typing import List, Optional, Sequence
import docx
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.shared import Inches

from .section_ast import SectionUnit


class DocxSerializer:
    """
    Deep serialization module converting SectionUnit lists into publisher-grade DOCX.
    """

    @staticmethod
    def render_sections_to_docx(
        sections: Sequence[SectionUnit],
        output_docx_path: str | Path,
        template_docx_path: Optional[str | Path] = None,
        base_dir: Optional[str | Path] = None,
    ) -> Path:
        """
        Renders a list of SectionUnits into a target DOCX file.
        Hides all table, image, heading, and style injection details.
        """
        out_p = Path(output_docx_path).resolve()
        out_p.parent.mkdir(parents=True, exist_ok=True)
        base_p = Path(base_dir).resolve() if base_dir else out_p.parent

        if template_docx_path and Path(template_docx_path).exists():
            doc = docx.Document(str(template_docx_path))
            # Clear template body paragraphs and tables while preserving styles.xml, fonts, geometry
            for p in list(doc.paragraphs):
                p._element.getparent().remove(p._element)
            for t in list(doc.tables):
                t._element.getparent().remove(t._element)
        else:
            doc = docx.Document()

        for sec in sections:
            if sec.title != "__PREAMBLE__":
                doc.add_paragraph(sec.title, style=f"Heading {min(sec.level, 3)}")

            paras = sec.effective_content.split("\n\n")
            for p_txt in paras:
                p_strip = p_txt.strip()
                if not p_strip:
                    continue

                # 1. Heading match
                m_h = re.match(r'^(#+)\s+(.+)$', p_strip)
                if m_h:
                    h_txt = m_h.group(2).strip()
                    h_txt = re.sub(r'\{custom-style="[^"]+"\}', '', h_txt).strip()
                    if sec.title != "__PREAMBLE__" and h_txt == sec.title:
                        continue
                    h_lvl = min(len(m_h.group(1)), 3)
                    doc.add_paragraph(h_txt, style=f"Heading {h_lvl}")
                    continue

                # 2. Image token match: ![](images/...) or ![alt](path){...}
                m_img = re.search(r'!\[.*?\]\((.*?)\)', p_strip)
                if m_img:
                    img_rel = m_img.group(1).split('{')[0].strip()
                    img_full = (base_p / img_rel).resolve()
                    if img_full.exists():
                        p_img = doc.add_paragraph()
                        try:
                            p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                        except Exception:
                            pass
                        run = p_img.add_run()
                        try:
                            run.add_picture(str(img_full), width=Inches(4.5))
                            continue
                        except Exception:
                            pass

                # 3. Table match: markdown table
                tbl_lines = [l.strip() for l in p_strip.splitlines() if l.strip()]
                if len(tbl_lines) >= 2 and tbl_lines[0].startswith('|') and tbl_lines[0].endswith('|') and re.match(r'^\|[\s\-:|]+\|$', tbl_lines[1]):
                    rows_data = []
                    for line in tbl_lines:
                        if re.match(r'^\|[\s\-:|]+\|$', line):
                            continue
                        cells = [c.strip() for c in line.strip('|').split('|')]
                        rows_data.append(cells)
                    if rows_data:
                        max_cols = max(len(r) for r in rows_data)
                        tbl = doc.add_table(rows=len(rows_data), cols=max_cols)
                        try:
                            tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                        except Exception:
                            pass
                        try:
                            tbl.style = 'Table Grid'
                        except Exception:
                            pass
                        for r_idx, r_data in enumerate(rows_data):
                            row = tbl.rows[r_idx]
                            for c_idx, cell_val in enumerate(r_data):
                                if c_idx < len(row.cells):
                                    row.cells[c_idx].text = cell_val
                        continue

                # 4. Custom-style block: ::: {custom-style="名稱"} ... :::
                m_st = re.search(r':::\s*\{\s*custom-style="([^"]+)"\s*\}\n?(.*?)\n?:::', p_strip, re.DOTALL)
                if m_st:
                    style_name = m_st.group(1)
                    body = m_st.group(2).strip()
                    try:
                        doc.add_paragraph(body, style=style_name)
                    except Exception:
                        doc.add_paragraph(body)
                else:
                    doc.add_paragraph(p_strip)

        doc.save(str(out_p))
        return out_p
