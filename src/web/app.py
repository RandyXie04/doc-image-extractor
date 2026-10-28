import os
import io
import sys
import time
import shutil
import uuid
import asyncio
import base64
import secrets
from contextlib import asynccontextmanager
from pathlib import Path

from fastapi import FastAPI, UploadFile, File, Form, BackgroundTasks, HTTPException, Request, Response
from fastapi.responses import HTMLResponse, FileResponse, StreamingResponse
from fastapi.staticfiles import StaticFiles
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel

import warnings
warnings.filterwarnings("ignore", message=".*The `fitz` API is deprecated.*")
import pymupdf as fitz
from PIL import Image, ImageDraw

from src.core_agent import PDFConversionAgent
from config import PATHS, CFG
from src.scripts.cleanup_scratch import cleanup_scratch
from src.utils.file_guard import (
    MaxBodySizeMiddleware,
    save_upload_file_safely,
    purge_expired_files,
    purge_in_memory_records,
)

# --- Pydantic Data Models / Schemas ---
class MarkdownUpdate(BaseModel):
    content: str

def run_full_cleanup() -> dict:
    """
    全系統 24 小時無日誌與過期檔案安全清理器 (Full 24-Hour Janitor)
    - 安全銷毀超期 (>24h) 之 03_output (含 backup_originals)、01_input、database_text、scratch、系統 Temp
    - 保護正在執行之活躍任務檔案 (active_paths)，避免 Windows WinError 32 異常
    - 同步清除 SQLite task_records 與 In-Memory 狀態字典中超時紀錄
    - 嚴格遵守無日誌架構 (Zero-Log)：若非必要除錯，清空/銷毀磁碟日誌
    """
    import time
    from src.utils.task_queue import get_task_queue

    max_age_seconds = CFG.file_retention_hours * 3600
    queue = get_task_queue()

    # 1. 收集正在處理中的檔案路徑 (白名單保護，避免 Windows 鎖定或損毀處理中檔案)
    active_paths: set[Path] = set()
    try:
        active_paths.update(queue.get_active_payload_paths())
    except Exception:
        pass

    for tid, tinfo in list(tasks.items()):
        if isinstance(tinfo, dict) and tinfo.get("status") in ("queued", "processing"):
            for key in ("result_file", "word_file", "zip_file"):
                val = tinfo.get(key)
                if val:
                    active_paths.add(Path(val))

    # 2. 定義所有需要巡檢銷毀之目標資料夾
    target_dirs = [
        PATHS.output_dir,                                  # data/03_output (含 backup_originals)
        PATHS.input_dir,                                   # data/01_input
        PATHS.data_dir / "database_text",                  # data/database_text (含 cropped_*.pdf)
        PATHS.scratch_dir,                                 # scratch/
        PATHS.temp_base,                                   # %TEMP%/doc-image-extractor-temp (含 ai_heading_cache, ai_kaiti_cache, previews 等)
        PATHS.logs_dir,                                    # logs/
    ]

    # 3. 執行磁碟檔案銷毀 (跳過受保護的範本檔案)
    deleted_files_count = purge_expired_files(
        target_dirs=target_dirs,
        max_age_seconds=max_age_seconds,
        protected_names={".gitkeep", "README.md", "user_template.docx", "template.docx"},
        active_paths=active_paths,
    )

    # 4. 同步清除 SQLite task_queue 超過 24 小時之歷史任務
    purged_tasks_count = 0
    try:
        purged_tasks_count = queue.cleanup_expired_tasks(max_age_seconds=max_age_seconds)
    except Exception:
        pass

    # 5. 同步淘汰 In-Memory 字典項目，杜絕伺服器長久運作內存洩漏
    active_keys = {
        k for k, v in tasks.items() if isinstance(v, dict) and v.get("status") in ("queued", "processing")
    }
    purged_mem_tasks = purge_in_memory_records(tasks, max_age_seconds=max_age_seconds, active_keys=active_keys)

    active_ocr_stems = {
        stem for stem, v in ocr_progress_dict.items() if isinstance(v, dict) and v.get("status") == "processing"
    }
    purged_mem_ocr = purge_in_memory_records(ocr_progress_dict, max_age_seconds=max_age_seconds, active_keys=active_ocr_stems)
    purged_mem_img = purge_in_memory_records(extracted_image_tasks, max_age_seconds=max_age_seconds)

    # 6. 無日誌規範：若未啟用磁碟日誌，清理 logs 目錄現存殘留
    if not CFG.enable_disk_logging and PATHS.logs_dir.exists():
        for log_file in PATHS.logs_dir.glob("*.log"):
            try:
                log_file.unlink(missing_ok=True)
            except OSError:
                pass

    return {
        "deleted_files": deleted_files_count,
        "purged_db_tasks": purged_tasks_count,
        "purged_mem_records": purged_mem_tasks + purged_mem_ocr + purged_mem_img,
    }

# 向下相容既有模組與測試之別名
cleanup_old_files = run_full_cleanup

@asynccontextmanager
async def lifespan(app_instance: FastAPI):
    # 確保所有需要的目錄已建立
    PATHS.ensure_all()
    
    # 首次啟動在背景自動偵測並準備模型，避免阻塞伺服器啟動導致 WebView 載入超時
    try:
        import threading
        from src.scripts.model_manager import ensure_model_ready
        threading.Thread(target=ensure_model_ready, daemon=True, name="ModelInitThread").start()
    except Exception as e:
        print(f"[Startup Warning] Model background preparation failed: {e}")

    # 每月 1 號自動清空 scratch 暫存
    try:
        cleanup_scratch(force=False, log_fn=print)
    except Exception as e:
        print(f"[Startup Warning] Scratch cleanup failed: {e}")

    # 啟動時立即執行全系統 24 小時過期清理
    try:
        stats = run_full_cleanup()
        if stats["deleted_files"] > 0 or stats["purged_db_tasks"] > 0:
            print(f"[Janitor] 啟動清理完成：銷毀 {stats['deleted_files']} 個過期檔案，清理 {stats['purged_db_tasks']} 筆過期任務")
    except Exception as e:
        print(f"[Startup Warning] Old files cleanup failed: {e}")

    # 啟動時自動檢查並修復異常中斷之佇列任務
    try:
        from src.utils.task_queue import get_task_queue
        recovered_count = get_task_queue().recover_interrupted_tasks()
        if recovered_count > 0:
            print(f"[TaskQueue] 已自動重置/標記 {recovered_count} 個中斷任務")
    except Exception as e:
        print(f"[Startup Warning] Task queue recovery failed: {e}")

    # 啟動常態定時清理背景守護程序 (每 30 分鐘巡檢一次)
    async def periodic_cleanup_daemon():
        while True:
            try:
                await asyncio.sleep(1800)
                run_full_cleanup()
            except asyncio.CancelledError:
                break
            except Exception as e:
                print(f"[Janitor Warning] Periodic cleanup error: {e}")

    cleanup_daemon_task = asyncio.create_task(periodic_cleanup_daemon())
    _BACKGROUND_TASKS.add(cleanup_daemon_task)
    cleanup_daemon_task.add_done_callback(_BACKGROUND_TASKS.discard)

    try:
        yield
    finally:
        cleanup_daemon_task.cancel()

app = FastAPI(title="PDF AI 公式萃取站", lifespan=lifespan)

# 限制重度 OCR / 版面模型之併發執行數為 1，防止區網多人同時上傳爆 VRAM / RAM
_OCR_SEMAPHORE = asyncio.Semaphore(1)
# 保留背景 Task 參照，防止 Python 3.10+ 在長任務執行期間因 GC 導致 Task 丟失
_BACKGROUND_TASKS: set[asyncio.Task] = set()

# 第一道防線：在 Starlette 解析 Multipart 寫入 SpooledTemporaryFile 前檢查 Content-Length
app.add_middleware(
    MaxBodySizeMiddleware,
    max_upload_bytes=CFG.max_upload_size_mb * 1024 * 1024,
    max_template_bytes=CFG.max_template_size_mb * 1024 * 1024,
)

@app.middleware("http")
async def add_no_cache_header(request: Request, call_next):
    response = await call_next(request)
    if request.url.path.startswith("/static/"):
        response.headers["Cache-Control"] = "no-cache, no-store, must-revalidate"
        response.headers["Pragma"] = "no-cache"
        response.headers["Expires"] = "0"
    return response

# --- 鑑權層防護 (Authentication Guard Middleware) ---
@app.middleware("http")
async def authentication_guard_middleware(request: Request, call_next):
    """
    鑑權層：當 CFG.auth_enabled 為 True 時，強制校驗 HTTP Basic 憑證或 X-API-Key Header。
    允許 CORS Preflight (OPTIONS) 通行以確保協定相容性。
    """
    if not CFG.auth_enabled:
        return await call_next(request)

    # 允許 CORS preflight
    if request.method == "OPTIONS":
        return await call_next(request)

    authenticated = False
    auth_header = request.headers.get("Authorization")
    api_key = request.headers.get("X-API-Key")

    # 1. 支援 X-API-Key Header (供自動化呼叫或 API 存取)
    if api_key and CFG.auth_password and secrets.compare_digest(api_key, CFG.auth_password):
        authenticated = True

    # 2. 支援 HTTP Basic Auth (使用者:密碼，供瀏覽器原生對話框)
    elif auth_header and auth_header.startswith("Basic "):
        try:
            encoded_creds = auth_header.split(" ", 1)[1]
            decoded = base64.b64decode(encoded_creds).decode("utf-8")
            if ":" in decoded:
                username, password = decoded.split(":", 1)
                user_match = secrets.compare_digest(username, CFG.auth_username)
                pass_match = secrets.compare_digest(password, CFG.auth_password)
                if user_match and pass_match:
                    authenticated = True
        except Exception:
            authenticated = False

    if not authenticated:
        return Response(
            content="Authentication required: 請提供正確的 HTTP Basic 憑證或 X-API-Key Header",
            status_code=401,
            headers={"WWW-Authenticate": 'Basic realm="Doc-Image-Extractor"'},
            media_type="text/plain; charset=utf-8",
        )

    return await call_next(request)


# --- 協定層防護 (Protocol / CORS Middleware) ---
# 移除非法之 allow_origins=["*"] 搭配 allow_credentials=True；限制僅允許 Localhost 與私有區域網路 (RFC 1918)
cors_origins = CFG.allowed_origins if CFG.allowed_origins else []
app.add_middleware(
    CORSMiddleware,
    allow_origins=cors_origins,
    allow_origin_regex=r"^https?://(localhost|127\.0\.0\.1|192\.168\.\d+\.\d+|10\.\d+\.\d+\.\d+|172\.(1[6-9]|2\d|3[01])\.\d+\.\d+)(:\d+)?$",
    allow_credentials=True,
    allow_methods=["GET", "POST", "PUT", "DELETE", "OPTIONS"],
    allow_headers=["*"],
)

# Simple in-memory task tracker
tasks = {}
# Global OCR progress tracking
ocr_progress_dict = {}
# Extracted images task tracking
extracted_image_tasks = {}

@app.post("/api/upload_file")
async def upload_file(file: UploadFile = File(...)):
    orig_ext = Path(file.filename).suffix.lower() if file.filename else ""
    if orig_ext != ".pdf":
        raise HTTPException(status_code=400, detail="公式萃取與預覽僅支援 .pdf 格式檔案")
        
    file_id = f"{uuid.uuid4()}.pdf"
    file_path = PATHS.input_dir / file_id
    
    # 第二道防線：安全串流分塊寫入，若累積超標立即中止並刪除半成品檔案
    await save_upload_file_safely(file, file_path, max_mb=CFG.max_upload_size_mb)
        
    total_pages = 1
    try:
        with fitz.open(file_path) as doc:
            total_pages = len(doc)
    except Exception as e:
        if file_path.exists():
            file_path.unlink(missing_ok=True)
        raise HTTPException(status_code=400, detail=f"無法解析此 PDF 檔案: {e}")

    orig_name = Path(file.filename).stem if file.filename else "已轉檔"
    return {"file_id": file_id, "total_pages": total_pages, "orig_name": orig_name}

@app.post("/api/upload_template")
async def upload_template(file: UploadFile = File(...)):
    if not file.filename or not file.filename.endswith('.docx'):
        raise HTTPException(status_code=400, detail="僅支援上傳 .docx 格式的範本檔案")
    
    template_dir = PATHS.data_dir / "database_text" / "custom_templates"
    template_dir.mkdir(parents=True, exist_ok=True)
    
    file_path = template_dir / "user_template.docx"
    await save_upload_file_safely(file, file_path, max_mb=CFG.max_template_size_mb)
        
    return {"status": "success", "message": f"範本 {file.filename} 上傳成功", "filename": file.filename}

@app.get("/api/get_current_template")
async def get_current_template():
    template_path = PATHS.data_dir / "database_text" / "custom_templates" / "user_template.docx"
    default_template_path = PATHS.data_dir / "database_text" / "template.docx"
    from src.utils.path_helper import get_template_path
    bundled_template_path = get_template_path()
    
    if template_path.exists():
        # Ideally we might want to store the original filename in a metadata file, 
        # but for simplicity we just return a static name or checking existence.
        return {"has_custom": True, "name": "user_template.docx"}
    elif default_template_path.exists():
        return {"has_custom": False, "name": "系統自訂範本"}
    elif bundled_template_path.exists():
        return {"has_custom": False, "name": "系統預設範本"}
    else:
        return {"has_custom": False, "name": "無可用範本"}

@app.get("/api/get_template_styles")
async def get_template_styles():
    import zipfile
    import xml.etree.ElementTree as ET
    from src.utils.path_helper import get_template_path
    
    template_path = PATHS.data_dir / "database_text" / "custom_templates" / "user_template.docx"
    if not template_path.exists():
        template_path = PATHS.data_dir / "database_text" / "template.docx"
        if not template_path.exists():
            template_path = get_template_path()
            if not template_path.exists():
                return {"status": "error", "message": "無可用範本"}
            
    styles = []
    W_NS = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"
    try:
        with zipfile.ZipFile(template_path, 'r') as z:
            if 'word/styles.xml' in z.namelist():
                tree = ET.fromstring(z.read('word/styles.xml'))
                for s in tree.findall(f"{W_NS}style"):
                    if s.get(f"{W_NS}type") == "paragraph":
                        name_el = s.find(f"{W_NS}name")
                        if name_el is not None:
                            name_val = name_el.get(f"{W_NS}val")
                            if name_val:
                                styles.append(name_val)
    except Exception as e:
        return {"status": "error", "message": f"解析樣式失敗: {e}"}
        
    return {"status": "success", "styles": styles}

@app.get("/api/version")
async def get_version():
    from config import VERSION
    return {
        "app_version": VERSION.app_version,
        "model_version": VERSION.model_version
    }

@app.post("/api/check_update")
async def check_update():
    from config import VERSION
    from src.scripts.updater_service import check_latest_release
    info = check_latest_release()
    if not info:
        return {"has_update": False, "status": "error", "msg": "無法連線至 GitHub 檢查更新。"}
    
    if info.get('status') == 'no_releases':
        info['has_update'] = False
        info['msg'] = "目前線上尚未發布新版本，您使用的是最新本機版本。"
        return info

    if info.get('status') == 'error':
        return info
    
    try:
        if info.get('latest_version'):
            remote_ver = [int(x) for x in info['latest_version'].split('.')]
            local_ver = [int(x) for x in VERSION.app_version.split('.')]
            has_update = remote_ver > local_ver
        else:
            has_update = False
    except Exception:
        has_update = False
            
    info['has_update'] = has_update
    return info

@app.post("/api/apply_update")
async def apply_update(request: Request):
    # 1. 僅允許本機操作更新，防止區網未授權遠端觸發更新或重啟
    client_host = request.client.host if request.client else ""
    if client_host not in ("127.0.0.1", "::1", "localhost", "testclient"):
        raise HTTPException(status_code=403, detail="安全限制：系統更新僅允許伺服器本機操作")

    data = await request.json()
    download_url = data.get('download_url')
    if not download_url:
        return {"status": "error", "msg": "No download URL provided."}

    # 2. 嚴格驗證更新來源網址
    import urllib.parse
    parsed = urllib.parse.urlparse(download_url)
    if parsed.scheme != 'https' or parsed.netloc not in ('github.com', 'objects.githubusercontent.com'):
        raise HTTPException(status_code=400, detail="不合法的更新下載來源網址")
        
    from src.scripts.updater_service import perform_update
    import threading
    import time
    
    result = perform_update(download_url)
    if result.get('status') == 'success':
        def delayed_exit():
            time.sleep(1)
            os._exit(0)
        threading.Thread(target=delayed_exit, daemon=True).start()
    return result

@app.get("/api/render_preview/{file_id}")
async def render_preview(file_id: str, page: int = 1, header: float = 0.1, footer: float = 0.1, left: float = 0.0, right: float = 0.0):
    """回傳帶有裁切輔助線（上下紅藍、左右綠）的單頁 PDF 預覽圖"""
    file_path = PATHS.input_dir / file_id
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
        
    try:
        with fitz.open(file_path) as doc:
            page_idx = max(0, min(page - 1, len(doc) - 1))
            page_obj: Any = doc[page_idx]
            
            # 渲染成圖片
            pix = page_obj.get_pixmap(dpi=72)
            img = Image.frombytes("RGB", (pix.width, pix.height), pix.samples)
            
            # 繪製輔助線
            draw = ImageDraw.Draw(img)
            w, h = img.size
            y_header = int(h * header)
            y_footer = int(h * (1.0 - footer))
            x_left = int(w * left)
            x_right = int(w * (1.0 - right))
            
            # 上下水平輔助線 (頂部紅、底部藍)
            draw.line([(0, y_header), (w, y_header)], fill="red", width=2)
            draw.line([(0, y_footer), (w, y_footer)], fill="blue", width=2)

            # 左右垂直輔助線 (綠色)
            if x_left > 0:
                draw.line([(x_left, 0), (x_left, h)], fill="green", width=2)
            if right > 0:
                draw.line([(x_right, 0), (x_right, h)], fill="green", width=2)
            
            buf = io.BytesIO()
            img.save(buf, format="JPEG", quality=80)
            buf.seek(0)
            
            return StreamingResponse(buf, media_type="image/jpeg")
    except Exception as e:
        raise HTTPException(status_code=500, detail=str(e))

def process_pdf(task_id: str, file_path: str, extract_formulas: bool, start_page: int, end_page: int, header_ratio: float, footer_ratio: float, left_ratio: float = 0.0, right_ratio: float = 0.0, extract_inline: bool = False):
    try:
        tasks[task_id]["status"] = "processing"
        tasks[task_id]["progress"] = 5.0
        
        def progress_cb(current, total, msg):
            pct = 5.0 + (current / total) * 90.0
            tasks[task_id]["progress"] = pct
            tasks[task_id]["message"] = msg
            tasks[task_id]["log"] += f"[{current}/{total}] {msg}\n"
            
        def log_fn(msg):
            tasks[task_id]["log"] += f"{msg}\n"
            
        log_fn(f"[DEBUG] 轉檔配置: extract_formulas={extract_formulas}, extract_inline={extract_inline}")

        agent = PDFConversionAgent(
            input_pdf=file_path, 
            header_ratio=header_ratio, 
            footer_ratio=footer_ratio,
            left_ratio=left_ratio,
            right_ratio=right_ratio,
            extract_inline=extract_inline
        )
        
        pipeline_res = agent.execute_pipeline(
            extract_formulas=extract_formulas,
            start_page_idx=start_page,
            end_page_idx=end_page if end_page > 0 else None,
            log_fn=log_fn,
            progress_callback=progress_cb
        )
        
        if pipeline_res and pipeline_res.get("delivery_folder"):
            tasks[task_id]["word_file"] = None
            tasks[task_id]["zip_file"] = pipeline_res.get("zip_path")
            tasks[task_id]["result_file"] = pipeline_res.get("zip_path")
            tasks[task_id]["progress"] = 100.0
            tasks[task_id]["status"] = "completed"
            tasks[task_id]["message"] = "處理完成"
        else:
            raise Exception("未能產生有效輸出成果")
            
    except Exception as e:
        tasks[task_id]["status"] = "failed"
        tasks[task_id]["message"] = f"錯誤: {str(e)}"
        tasks[task_id]["log"] += f"\n發生異常：{str(e)}\n"

@app.post("/api/process")
async def start_process(
    background_tasks: BackgroundTasks,
    file_id: str = Form(...),
    extract_formulas: bool = Form(True),
    start_page: int = Form(0),
    end_page: int = Form(0),
    header_ratio: float = Form(0.1),
    footer_ratio: float = Form(0.1),
    left_ratio: float = Form(0.0),
    right_ratio: float = Form(0.0),
    extract_inline: bool = Form(False),
    orig_name: str = Form("")
):
    file_path = PATHS.input_dir / file_id
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="File not found")
        
    task_id = str(uuid.uuid4())
    tasks[task_id] = {
        "status": "queued",
        "progress": 0.0,
        "message": "排隊中...",
        "log": "任務已加入佇列...\n",
        "result_file": None,
        "word_file": None,
        "zip_file": None,
        "orig_name": orig_name or "已轉檔",
        "created_at": time.time(),
    }
    
    background_tasks.add_task(
        process_pdf, 
        task_id, 
        str(file_path), 
        extract_formulas, 
        start_page, 
        end_page, 
        header_ratio, 
        footer_ratio,
        left_ratio,
        right_ratio,
        extract_inline
    )
    return {"task_id": task_id}

@app.get("/api/app_mode")
async def get_app_mode():
    """取得當前運行模式：editor (出版社編輯) 或 dev (內部工程師)"""
    from config import CFG
    return {"mode": CFG.app_mode}

@app.post("/api/set_app_mode")
async def set_app_mode(request: Request):
    """動態切換運行模式 (免重啟)"""
    data = await request.json()
    new_mode = data.get("mode", "editor").lower()
    if new_mode not in ["editor", "dev"]:
        raise HTTPException(status_code=400, detail="模式必須為 'editor' 或 'dev'")
    import config.settings
    object.__setattr__(config.settings.CFG, "app_mode", new_mode)
    return {"status": "success", "mode": new_mode}



# Mount static files
from src.utils.path_helper import get_static_dir
static_dir = get_static_dir()

os.makedirs(static_dir, exist_ok=True)
app.mount("/static", StaticFiles(directory=static_dir), name="static")

@app.get("/favicon.ico", include_in_schema=False)
async def favicon():
    svg_icon = (
        '<svg xmlns="http://www.w3.org/2000/svg" viewBox="0 0 24 24" fill="none" stroke="#ea580c" stroke-width="2" stroke-linecap="round" stroke-linejoin="round">'
        '<path d="M4 19.5v-15A2.5 2.5 0 0 1 6.5 2H20v20H6.5a2.5 2.5 0 0 1-2.5-2.5Z"/>'
        '<path d="M6 6h10"/>'
        '<path d="M6 10h10"/>'
        '</svg>'
    )
    return Response(content=svg_icon, media_type="image/svg+xml")

NO_CACHE_HEADERS = {
    "Cache-Control": "no-cache, no-store, must-revalidate",
    "Pragma": "no-cache",
    "Expires": "0"
}

@app.get("/", response_class=HTMLResponse)
async def read_index():
    index_path = static_dir / "index.html"
    if index_path.exists():
        with open(index_path, "r", encoding="utf-8") as f:
            return HTMLResponse(content=f.read(), headers=NO_CACHE_HEADERS)
    raise HTTPException(status_code=404, detail="Index page not found")

@app.get("/{page}.html", response_class=HTMLResponse)
async def read_page(page: str):
    file_path = static_dir / f"{page}.html"
    if file_path.exists():
        with open(file_path, "r", encoding="utf-8") as f:
            return HTMLResponse(content=f.read(), headers=NO_CACHE_HEADERS)
    raise HTTPException(status_code=404, detail="Page not found")

@app.get("/api/status/{task_id}")
async def get_status(task_id: str):
    if task_id not in tasks:
        return {"status": "not_found"}
    return tasks[task_id]

@app.get("/api/system_config")
async def get_system_config():
    """取得系統安全防護與無日誌配置資訊"""
    from config import CFG
    return {
        "max_upload_size_mb": CFG.max_upload_size_mb,
        "max_template_size_mb": CFG.max_template_size_mb,
        "file_retention_hours": CFG.file_retention_hours,
        "enable_disk_logging": CFG.enable_disk_logging,
        "app_mode": CFG.app_mode
    }

@app.get("/api/download/{task_id}")
async def download_result(task_id: str):
    """向下相容通用下載端點"""
    if task_id in tasks and tasks[task_id].get("result_file"):
        path = tasks[task_id]["result_file"]
        if not os.path.exists(path):
            raise HTTPException(status_code=404, detail="檔案不存在或已超過 24 小時保存期限由系統自動銷毀")
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document" if path.endswith(".docx") else "application/zip"
        return FileResponse(path, media_type=media_type, filename=os.path.basename(path))
    raise HTTPException(status_code=404, detail="成果檔案不存在或任務未完成")

@app.get("/api/hardware_status")
async def get_hardware_status():
    from config import CFG
    profile = CFG.get_hardware_profile()
    return profile

@app.get("/api/download/{task_id}/formulas")
async def download_formulas(task_id: str):
    """專用公式圖檔包下載端點"""
    if task_id in tasks and tasks[task_id].get("zip_file"):
        path = tasks[task_id]["zip_file"]
        if not os.path.exists(path):
            raise HTTPException(status_code=404, detail="檔案不存在或已超過 24 小時保存期限由系統自動銷毀")
        return FileResponse(
            path,
            media_type="application/zip",
            filename=os.path.basename(path)
        )
    raise HTTPException(status_code=404, detail="公式圖檔包不存在或任務未完成")

# =========================================================================
# 文件圖片無損提取 API (DOCX / PDF)
# =========================================================================
@app.post("/api/extract_images")
async def api_extract_images(
    file: UploadFile = File(...),
    to_grayscale: bool = Form(False)
):
    from src.scripts.extract_images import process_document_images
    import asyncio
    import time

    orig_filename = file.filename or "document"
    orig_ext = Path(orig_filename).suffix.lower()
    if orig_ext not in [".docx", ".pdf"]:
        raise HTTPException(status_code=400, detail="僅支援 .docx 與 .pdf 檔案格式")

    task_id = str(uuid.uuid4())
    temp_dir = PATHS.root / "scratch" / f"extract_{task_id}"
    input_file_path = temp_dir / orig_filename
    extracted_folder = temp_dir / "images"
    output_zip_name = f"{Path(orig_filename).stem}_extracted_images.zip"
    output_zip_path = PATHS.root / "data" / "03_output" / output_zip_name

    temp_dir.mkdir(parents=True, exist_ok=True)
    await save_upload_file_safely(file, input_file_path, max_mb=CFG.max_upload_size_mb)

    try:
        count, zip_path = await asyncio.to_thread(
            process_document_images,
            file_path=str(input_file_path),
            output_zip_path=str(output_zip_path),
            temp_dir=str(extracted_folder),
            to_grayscale=to_grayscale
        )

        # 雙重驗證：檢查實際在資料夾中產生的檔案數量
        actual_files = 0
        if extracted_folder.exists():
            actual_files = len([f for f in extracted_folder.iterdir() if f.is_file()])
        
        if actual_files == 0:
            count = 0
            
        if count == 0:
            return {
                "success": False,
                "count": 0,
                "message": "未在此文件中偵測到任何內嵌圖片，或提取過程未產出任何圖檔。"
            }

        extracted_image_tasks[task_id] = {
            "zip_path": zip_path,
            "filename": output_zip_name,
            "count": actual_files,
            "created_at": time.time(),
        }

        return {
            "success": True,
            "count": actual_files,
            "download_url": f"/api/download_extracted_images/{task_id}",
            "filename": output_zip_name,
            "source_path": f"data/03_output/{output_zip_name}",
            "message": f"成功提取 {actual_files} 張圖片！"
        }
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"圖片提取失敗: {str(e)}")
    finally:
        try:
            shutil.rmtree(temp_dir, ignore_errors=True)
        except Exception:
            pass

@app.get("/api/download_extracted_images/{task_id}")
async def download_extracted_images(task_id: str):
    if task_id in extracted_image_tasks:
        info = extracted_image_tasks[task_id]
        zip_path = info["zip_path"]
        if os.path.exists(zip_path):
            return FileResponse(
                zip_path,
                media_type="application/zip",
                filename=info["filename"]
            )
        raise HTTPException(status_code=404, detail="檔案不存在或已超過 24 小時保存期限由系統自動銷毀")
    raise HTTPException(status_code=404, detail="找不到提取的壓縮檔案或任務不存在")


@app.post("/api/open_folder/{task_id}")
async def open_folder(task_id: str, request: Request):
    """在 Windows 檔案總管中開啟成果所在資料夾並選取檔案"""
    client_host = request.client.host if request.client else ""
    if client_host not in ("127.0.0.1", "::1", "localhost", "testclient"):
        return {"status": "fallback", "message": "伺服器本機資料夾僅限本機瀏覽，遠端區網連線請使用瀏覽器下載"}

    if task_id in tasks:
        target = tasks[task_id].get("zip_file") or tasks[task_id].get("result_file")
        if target and os.path.exists(target):
            import subprocess
            subprocess.Popen(f'explorer /select,"{os.path.abspath(target)}"')
            return {"status": "success"}
    raise HTTPException(status_code=404, detail="成果檔案不存在或尚未生成")




# =========================================================================
# Founder Tools API (方正排版修復)
# =========================================================================
def _cleanup_temp_files(*file_paths):
    """背景清理暫存檔，防止磁碟洩漏"""
    for p in file_paths:
        if p and os.path.exists(p):
            try:
                os.unlink(p)
            except Exception:
                pass

@app.post("/api/founder/repair")
async def api_founder_repair(
    file: UploadFile = File(...)
):
    import os
    import shutil
    import asyncio
    import uuid
    from src.founder_tools.fix_founder_fonts import repair_pdf_file, repair_docx_file
    from config import PATHS
    
    if not file.filename:
        raise HTTPException(status_code=400, detail="未提供有效檔案名稱")

    filename = file.filename
    ext = os.path.splitext(filename)[1].lower()
    if ext not in [".pdf", ".docx"]:
        raise HTTPException(status_code=400, detail="不支援的檔案格式，請上傳 PDF 或 DOCX")

    task_id = str(uuid.uuid4())
    temp_input = PATHS.root / "scratch" / f"founder_repair_{task_id}{ext}"
    await save_upload_file_safely(file, temp_input, max_mb=CFG.max_upload_size_mb)
    created_temps = [str(temp_input)]
    
    output_dir = PATHS.root / "data" / "03_output"
    output_dir.mkdir(parents=True, exist_ok=True)

    try:
        if ext == ".pdf":
            temp_out_txt = str(PATHS.root / "scratch" / f"founder_repair_{task_id}.txt")
            output_docx = output_dir / f"repaired_{Path(filename).stem}_{task_id[:8]}.docx"
            created_temps.append(temp_out_txt)
            
            await asyncio.to_thread(repair_pdf_file, str(temp_input), temp_out_txt, str(output_docx))
            _cleanup_temp_files(*created_temps)
            
            return {
                "success": True,
                "source_path": f"data/03_output/{output_docx.name}",
                "filename": output_docx.name
            }
            
        elif ext == ".docx":
            output_docx = output_dir / f"repaired_{Path(filename).stem}_{task_id[:8]}.docx"
            
            await asyncio.to_thread(repair_docx_file, str(temp_input), str(output_docx))
            _cleanup_temp_files(*created_temps)
            
            return {
                "success": True,
                "source_path": f"data/03_output/{output_docx.name}",
                "filename": output_docx.name
            }
            
    except Exception as e:
        _cleanup_temp_files(*created_temps)
        raise HTTPException(status_code=500, detail=f"修復過程發生異常: {str(e)}")


@app.post("/api/report_issue")
async def report_issue(request: Request):
    import platform
    import urllib.parse
    
    data = await request.json()
    description = (data.get("description") or "").strip()[:2000]
    
    sys_info = f"OS: {platform.system()} {platform.release()}"
    body = f"**使用者回報:**\n{description}\n\n---\n**自動收集資訊:**\n```text\n{sys_info}\n```"
    
    repo = "RandyXie04/doc-image-extractor"
    encoded_title = urllib.parse.quote(f"[問題回報] {description[:30]}..." if description else "[問題回報] 請簡述問題")
    encoded_body = urllib.parse.quote(body)
    web_url = f"https://github.com/{repo}/issues/new?title={encoded_title}&body={encoded_body}"
    
    return {
        "status": "success",
        "url": web_url,
        "msg": "已生成 GitHub 官方回報頁面網址。"
    }





# =========================================================================
# 字體分析與樣式對應 API
# =========================================================================

@app.post("/api/scan_fonts")
async def api_scan_fonts(request: Request):
    """
    掃描指定 PDF 的字體清單及 template.docx 中的段落樣式。
    前端在「字體審閱」Modal 開啟前呼叫此路由。
    Body JSON: { "file_id": "...", "filename_orig": "..." }
    """
    import asyncio
    from src.scripts.pdf_font_analyzer import analyze_pdf_for_ui
    from src.utils.path_helper import get_template_path

    data = await request.json()
    file_id: str = data.get("file_id", "")
    filename_orig: str = data.get("filename_orig", "")

    # 解析 PDF 路徑（支援 input_dir 的上傳快取，或 database_text 的已存 PDF）
    pdf_path = None
    if file_id:
        candidate = PATHS.input_dir / file_id
        if candidate.exists():
            pdf_path = str(candidate)

    # 若 input_dir 找不到，嘗試 database_text
    if not pdf_path and filename_orig:
        candidate2 = PATHS.root / "data" / "database_text" / filename_orig
        if Path(candidate2).exists():
            pdf_path = str(candidate2)

    if not pdf_path:
        raise HTTPException(status_code=404, detail="找不到對應的 PDF 檔案，請先上傳 PDF。")

    # 解析 template.docx 路徑（按優先順序尋找）
    template_candidates = [
        PATHS.data_dir / "database_text" / "custom_templates" / "user_template.docx",
        PATHS.data_dir / "database_text" / "template.docx",
        get_template_path(),
    ]
    template_path = None
    for tc in template_candidates:
        if Path(str(tc)).exists():
            template_path = str(tc)
            break

    fonts_json_path = str(PATHS.fonts_json_path)

    try:
        result = await asyncio.to_thread(
            analyze_pdf_for_ui,
            pdf_path,
            template_path,
            fonts_json_path,
        )
        return result
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"字體分析失敗: {str(e)}")


@app.post("/api/save_font_mapping")
async def api_save_font_mapping(request: Request):
    """
    儲存人工審閱後的字體 → 樣式對應表至 fonts.json。
    Body JSON: { "mapping": { "FontName": {"role": "H2", "docx_style": "標題 2"}, ... }, "remove_keys": ["FontName", ...] }
    """
    from src.scripts.pdf_font_analyzer import save_fonts_json

    data = await request.json()
    mapping_raw: dict = data.get("mapping", {})
    remove_keys: list = data.get("remove_keys", [])

    if not mapping_raw and not remove_keys:
        return {"status": "success", "saved": 0, "message": "無變更項目需要儲存"}

    # 標記每筆為人工設定
    mapping = {}
    for font_name, info in mapping_raw.items():
        mapping[font_name] = {
            "role": info.get("role"),
            "docx_style": info.get("docx_style"),
            "mapped_by": "user",
        }

    fonts_json_path = str(PATHS.fonts_json_path)
    try:
        save_fonts_json(fonts_json_path, mapping, remove_keys=remove_keys)
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"儲存 fonts.json 失敗: {str(e)}")

    return {
        "status": "success",
        "saved": len(mapping),
        "removed": len(remove_keys),
        "message": f"已儲存字體樣式對應設定（更新 {len(mapping)} 筆，重置 {len(remove_keys)} 筆）"
    }


@app.get("/api/ocr_progress/{filename_stem}")
async def get_ocr_progress(filename_stem: str):
    progress = ocr_progress_dict.get(filename_stem)
    if progress is None:
        try:
            from src.utils.task_queue import get_task_queue
            queue = get_task_queue()
            rec = queue.get_task_by_stem(filename_stem)
            if rec:
                progress = {
                    "task_id": rec.task_id,
                    "progress": rec.progress,
                    "message": rec.message,
                    "status": rec.status,
                    "error": rec.error
                }
        except Exception:
            pass
    if progress is None:
        progress = {"progress": 0, "message": "Waiting...", "status": "processing"}
    return progress


@app.post("/api/upload_and_run_ocr")
async def upload_and_run_ocr(
    background_tasks: BackgroundTasks, 
    file: UploadFile = File(None), 
    file_id: str = Form(None),
    filename_orig: str = Form(None),
    style_mapping: str = Form(None),
    engine: str = Form("auto"),
    header_ratio: float = Form(0.1),
    footer_ratio: float = Form(0.1),
    left_ratio: float = Form(0.0),
    right_ratio: float = Form(0.0),
    start_page: int = Form(0),
    end_page: int = Form(0),
    no_footnote: bool = Form(True),
    ai_kaiti: bool = Form(False),
    ai_kaiti_threshold: float = Form(0.65)
):
    import sys
    import shutil
    import os
    from config import PATHS
    from src.utils.task_queue import get_task_queue, TaskRecord, TaskStatus
    
    pdf_dir = PATHS.root / 'data' / 'database_text'
    pdf_dir.mkdir(parents=True, exist_ok=True)
    
    # Support pre-uploaded file_id or direct upload
    if file_id and (PATHS.input_dir / file_id).exists():
        raw_pdf_path = PATHS.input_dir / file_id
        filename_stem = filename_orig or Path(file_id).stem
    elif file and file.filename:
        clean_filename = Path(file.filename).name
        ext = os.path.splitext(clean_filename)[1].lower()
        if ext != ".pdf":
            raise HTTPException(status_code=400, detail="僅支援 PDF 檔案格式")
        raw_pdf_path = pdf_dir / clean_filename
        await save_upload_file_safely(file, raw_pdf_path, max_mb=CFG.max_upload_size_mb)
        filename_stem = Path(clean_filename).stem
    else:
        raise HTTPException(status_code=400, detail="請提供 PDF 檔案或檔案代碼")

    task_id = str(uuid.uuid4())
    # WYSIWYG pre-crop and page range selection using fitz
    input_pdf_path = raw_pdf_path
    if header_ratio > 0 or footer_ratio > 0 or left_ratio > 0 or right_ratio > 0 or start_page > 0 or end_page > 0:
        cropped_pdf_path = pdf_dir / f"cropped_{filename_stem}_{task_id[:8]}.pdf"
        try:
            with fitz.open(raw_pdf_path) as doc:
                total_pages = len(doc)
                s_page = max(1, start_page) if start_page > 0 else 1
                e_page = min(total_pages, end_page) if end_page > 0 else total_pages
                
                # Filter pages by only keeping the selected range
                if s_page > 1 or e_page < total_pages:
                    doc.select(range(s_page - 1, e_page))
                
                for page in doc:
                    if header_ratio > 0 or footer_ratio > 0 or left_ratio > 0 or right_ratio > 0:
                        rect = page.rect
                        y_top = rect.height * max(0.0, min(header_ratio, 0.49))
                        f_r = footer_ratio if footer_ratio < 0.5 else (1.0 - footer_ratio)
                        y_bottom = rect.height * (1.0 - max(0.0, min(f_r, 0.49)))
                        x_left = rect.width * max(0.0, min(left_ratio, 0.49))
                        r_r = right_ratio if right_ratio < 0.5 else (1.0 - right_ratio)
                        x_right = rect.width * (1.0 - max(0.0, min(r_r, 0.49)))
                        page.set_cropbox(fitz.Rect(x_left, y_top, x_right, y_bottom))
                doc.save(cropped_pdf_path)
            input_pdf_path = cropped_pdf_path
        except Exception as crop_err:
            print(f"[Warning] PDF pre-crop/page-range failed: {crop_err}")
            input_pdf_path = raw_pdf_path
    queue = get_task_queue()
    queue.enqueue(
        TaskRecord(
            task_id=task_id,
            stem=filename_stem,
            status=TaskStatus.QUEUED.value,
            progress=0,
            message="正在初始化任務並排入佇列...",
            engine=engine,
            payload={
                "file_path": str(input_pdf_path),
                "ai_kaiti": ai_kaiti,
                "ai_kaiti_threshold": ai_kaiti_threshold,
                "no_footnote": no_footnote
            }
        )
    )

    ocr_progress_dict[filename_stem] = {
        "task_id": task_id,
        "progress": 0,
        "message": "正在初始化任務...",
        "status": "processing",
        "created_at": time.time(),
    }

    async def run_scripts_async(tid, pdf_path_str, stem, style_map_json, eng, skip_footnotes, ai_kaiti_on, ai_kaiti_thresh):
        import asyncio
        import json as _json
        try:
            ocr_progress_dict[stem]["message"] = "已排入轉檔佇列，等待運算資源..."
            queue.update_task(tid, status=TaskStatus.RUNNING.value, message="已排入轉檔佇列，等待運算資源...")
            async with _OCR_SEMAPHORE:
                output_dir = PATHS.root / 'data' / '03_output'
                backup_dir = output_dir / 'backup_originals'
                output_dir.mkdir(parents=True, exist_ok=True)
                
                md_file = output_dir / f"{stem}.md"
                docx_file = output_dir / f"{stem}.docx"
                
                import time as _time
                for f_path in [md_file, docx_file]:
                    if f_path.exists():
                        backup_dir.mkdir(parents=True, exist_ok=True)
                        backup_path = backup_dir / f"{f_path.stem}_{int(_time.time())}_{tid[:6]}{f_path.suffix}"
                        try: shutil.move(str(f_path), str(backup_path))
                        except: pass

                cmd = [
                    sys.executable, "-m", "src.scripts.pdf_engine_dispatcher",
                    "--file", pdf_path_str,
                    "--output_dir", str(output_dir),  # absolute path
                    "--output_stem", stem,
                    "--engine", eng
                ]
                if skip_footnotes:
                    cmd.append("--no_footnote")
                if ai_kaiti_on:
                    cmd.extend(["--ai_kaiti", "--ai_kaiti_threshold", str(ai_kaiti_thresh)])
                if style_map_json:
                    cmd.extend(["--style_mapping", style_map_json])
                    
                # Use asyncio subprocess to avoid blocking the event loop
                proc = await asyncio.create_subprocess_exec(
                    *cmd,
                    stdout=asyncio.subprocess.PIPE,
                    stderr=asyncio.subprocess.STDOUT,
                    cwd=str(PATHS.root)
                )
                
                if proc.stdout is not None:
                    async for raw_line in proc.stdout:
                        try:
                            line = raw_line.decode("utf-8", errors="replace").strip()
                        except Exception:
                            line = ""
                        if line:
                            if line.startswith("{") and "progress" in line:
                                try:
                                    data = _json.loads(line)
                                    prog = int(data.get("progress", ocr_progress_dict[stem]["progress"]))
                                    msg = data.get("message", line)
                                    ocr_progress_dict[stem]["progress"] = prog
                                    ocr_progress_dict[stem]["message"] = msg
                                    queue.update_task(tid, progress=prog, message=msg)
                                except Exception:
                                    pass
                            else:
                                short_msg = line[:120]
                                ocr_progress_dict[stem]["message"] = short_msg
                                queue.update_task(tid, message=short_msg)
                
                await proc.wait()
                
                if proc.returncode != 0:
                    last_msg = ocr_progress_dict[stem].get("message", "")
                    fail_msg = f"轉檔程序執行失敗 (代碼: {proc.returncode})"
                    if last_msg and not last_msg.startswith("{") and "轉檔程序執行失敗" not in last_msg:
                        fail_msg = f"{fail_msg}: {last_msg}"
                    ocr_progress_dict[stem]["status"] = "failed"
                    ocr_progress_dict[stem]["message"] = fail_msg
                    queue.update_task(tid, status=TaskStatus.FAILED.value, message=fail_msg, error=f"Process exited with {proc.returncode}: {last_msg}")
                    return

                if not md_file.exists():
                    fail_msg = "轉檔程序異常結束：未生成 Markdown 輸出檔案"
                    ocr_progress_dict[stem]["status"] = "failed"
                    ocr_progress_dict[stem]["message"] = fail_msg
                    queue.update_task(tid, status=TaskStatus.FAILED.value, message=fail_msg, error=fail_msg)
                    return

                docx_msg = "正在套用樣式並產生 Word (.docx) 文件..."
                ocr_progress_dict[stem]["progress"] = 92
                ocr_progress_dict[stem]["message"] = docx_msg
                queue.update_task(tid, progress=92, message=docx_msg)
                
                md_proc = await asyncio.create_subprocess_exec(
                    sys.executable, "-m", "src.scripts.md_to_docx",
                    "--files", str(md_file),
                    cwd=str(PATHS.root),
                    stdout=asyncio.subprocess.PIPE,
                    stderr=asyncio.subprocess.PIPE
                )
                _, md_err = await md_proc.communicate()
                
                if md_proc.returncode != 0 or not docx_file.exists():
                    err_detail = md_err.decode(errors="replace").strip() if md_err else ""
                    fail_msg = f"Word (.docx) 文件產生失敗: {err_detail[:150]}" if err_detail else "Word (.docx) 文件產生失敗"
                    ocr_progress_dict[stem]["status"] = "failed"
                    ocr_progress_dict[stem]["message"] = fail_msg
                    queue.update_task(tid, status=TaskStatus.FAILED.value, message=fail_msg, error=fail_msg)
                    return

                ocr_progress_dict[stem]["progress"] = 100
                ocr_progress_dict[stem]["message"] = "轉檔完成！"
                ocr_progress_dict[stem]["status"] = "completed"
                queue.update_task(tid, status=TaskStatus.COMPLETED.value, progress=100, message="轉檔完成！")
        except Exception as err:
            err_msg = f"處理過程異常: {str(err)}"
            ocr_progress_dict[stem]["status"] = "failed"
            ocr_progress_dict[stem]["message"] = err_msg
            queue.update_task(tid, status=TaskStatus.FAILED.value, message=err_msg, error=str(err))
        finally:
            if "cropped_" in os.path.basename(pdf_path_str) and os.path.exists(pdf_path_str):
                try:
                    os.unlink(pdf_path_str)
                except Exception:
                    pass

    task = asyncio.create_task(run_scripts_async(task_id, str(input_pdf_path), filename_stem, style_mapping, engine, no_footnote, ai_kaiti, ai_kaiti_threshold))
    _BACKGROUND_TASKS.add(task)
    task.add_done_callback(_BACKGROUND_TASKS.discard)
    return {"message": "OCR 管道已成功啟動", "filename_stem": filename_stem, "task_id": task_id}


@app.get("/api/queue/tasks")
async def list_queue_tasks(limit: int = 50, status: str | None = None):
    """取得分散式/持久化任務佇列之任務清單"""
    from src.utils.task_queue import get_task_queue
    queue = get_task_queue()
    records = queue.list_tasks(limit=limit, status=status)
    return {"tasks": [r.to_dict() for r in records]}


@app.post("/api/queue/cancel/{task_id}")
async def cancel_queue_task(task_id: str):
    """取消佇列中的任務"""
    from src.utils.task_queue import get_task_queue
    queue = get_task_queue()
    success = queue.cancel_task(task_id)
    return {"success": success, "task_id": task_id}


@app.get("/api/download_ocr_result/{filename_stem}/{ext}")
async def download_ocr_result(filename_stem: str, ext: str):
    from config import PATHS
    if ext not in ["docx", "md"]:
        raise HTTPException(status_code=400, detail="僅支援下載 docx 或 md")
    output_dir = PATHS.root / 'data' / '03_output'
    file_path = output_dir / f"{filename_stem}.{ext}"
    if not file_path.exists():
        raise HTTPException(status_code=404, detail="檔案尚未生成或已超過 24 小時保存期限由系統自動銷毀")
    return FileResponse(file_path, filename=f"{filename_stem}.{ext}")


@app.get("/api/get_markdown/{filename_stem}")
def get_markdown(filename_stem: str):
    md_file = PATHS.root / 'data' / '03_output' / f"{filename_stem}.md"
    if not md_file.exists():
        raise HTTPException(status_code=404, detail="Markdown file not found")
    with open(md_file, "r", encoding="utf-8") as f:
        return {"content": f.read()}


@app.post("/api/update_and_rebuild_docx/{filename_stem}")
async def update_and_rebuild_docx(filename_stem: str, data: MarkdownUpdate):
    md_file = PATHS.root / 'data' / '03_output' / f"{filename_stem}.md"
    if not md_file.exists():
        raise HTTPException(status_code=404, detail="Markdown file not found")
        
    def write_file():
        with open(md_file, "w", encoding="utf-8") as f:
            f.write(data.content)
            
    await asyncio.to_thread(write_file)
        
    # Re-run md_to_docx.py
    cmd = [sys.executable, "-m", "src.scripts.md_to_docx", "--files", str(md_file)]
    proc = await asyncio.create_subprocess_exec(
        *cmd,
        cwd=str(PATHS.root)
    )
    await proc.wait()
    
    return {"status": "success", "message": "已成功更新大綱並重新生成 Word 文件！"}

class NativeSaveRequest(BaseModel):
    source_path: str
    suggested_filename: str = ""

@app.post("/api/native_save_file")
async def api_native_save_file(req: NativeSaveRequest, request: Request):
    """
    [出版社編輯專用] 呼叫 Windows 原生檔案總管『另存新檔』視窗，
    具備嚴格路徑白名單校驗 (Path Traversal 防禦 - NIST PR.DS / ISO 27001 A.8.28)
    與非同步子進程防阻塞 (NIST PR.IP)。
    """
    # 0. 阻擋區網遠端客戶端在伺服器主機上彈出原生視窗
    client_host = request.client.host if request.client else ""
    if client_host not in ("127.0.0.1", "::1", "localhost", "testclient"):
        return {"status": "fallback", "message": "遠端區網連線請直接使用瀏覽器下載"}

    import asyncio
    import json
    import sys
    from config import PATHS
    
    # 1. 嚴格路徑校驗 (Path Traversal 防禦 - 使用 is_relative_to 防範前綴繞過)
    allowed_dirs = [
        (PATHS.root / "data" / "03_output").resolve(),
        (PATHS.root / "data" / "02_intermediate").resolve(),
        PATHS.data_dir.resolve()
    ]
    
    target_path = Path(req.source_path)
    if not target_path.is_absolute():
        target_path = (PATHS.root / req.source_path).resolve()
    else:
        target_path = target_path.resolve()
        
    is_safe = any(target_path.is_relative_to(d) for d in allowed_dirs)
    if not is_safe:
        raise HTTPException(status_code=403, detail="存取受限：路徑超出允許之專案輸出目錄")
        
    if not target_path.exists() or not target_path.is_file():
        raise HTTPException(status_code=404, detail="指定之輸出檔案不存在")
        
    suggested_filename = req.suggested_filename or target_path.name
    
    # 2. 非同步子進程執行原生對話框 (非阻塞 Event Loop)
    cmd = [
        sys.executable,
        "-m", "src.scripts.native_dialog",
        str(target_path),
        suggested_filename
    ]
    
    proc = None
    try:
        proc = await asyncio.create_subprocess_exec(
            *cmd,
            stdout=asyncio.subprocess.PIPE,
            stderr=asyncio.subprocess.PIPE,
            cwd=str(PATHS.root)
        )
        stdout, stderr = await asyncio.wait_for(proc.communicate(), timeout=300.0)
        output_str = stdout.decode("utf-8", errors="replace").strip()
        
        if output_str:
            try:
                res_data = json.loads(output_str)
                return res_data
            except Exception:
                pass
                
        return {"status": "cancelled", "message": "對話框已關閉或未選取檔案"}
    except asyncio.TimeoutError:
        if proc is not None:
            try:
                proc.kill()
            except Exception:
                pass
        return {"status": "fallback", "message": "操作逾時，請改用瀏覽器直接下載"}
    except Exception as e:
        return {"status": "error", "message": f"原生另存失敗: {str(e)}"}

class OpenFileFolderRequest(BaseModel):
    file_path: str

@app.post("/api/open_file_folder")
def api_open_file_folder(req: OpenFileFolderRequest, request: Request):
    """在 Windows 檔案總管中定位並選取已儲存之檔案"""
    client_host = request.client.host if request.client else ""
    if client_host not in ("127.0.0.1", "::1", "localhost", "testclient"):
        return {"status": "fallback", "message": "伺服器本機資料夾僅限本機瀏覽，遠端區網連線請使用瀏覽器下載"}

    import subprocess
    from config import PATHS
    
    # 嚴格路徑校驗 (Path Traversal 防禦 - 使用 is_relative_to 防範前綴繞過)
    allowed_dirs = [
        (PATHS.root / "data").resolve(),
        PATHS.data_dir.resolve()
    ]
    
    target = Path(req.file_path).resolve()
    is_safe = any(target.is_relative_to(d) for d in allowed_dirs)
    
    if not is_safe:
        raise HTTPException(status_code=403, detail="存取受限：路徑超出允許之專案資料目錄")
        
    if not target.exists():
        raise HTTPException(status_code=404, detail="目標檔案不存在")
        
    subprocess.Popen(f'explorer /select,"{str(target)}"')
    return {"status": "success"}


# --- POD Automated Markup Production API Endpoints ---
from typing import Optional, List, Dict, Any
from src.pod_studio.pod_service import (
    extract_raw_from_pdf,
    extract_raw_from_markdown,
    compile_blueprint_from_seed,
    propagate_and_validate_session,
)
from src.pod_studio.blueprint_schema import BookBlueprint
from src.pod_studio.vlm_continuation_engine import SecureVLMContinuationEngine



from src.web.schemas import (
    PodSectionPreview,
    PodExtractRawResponse,
    PodCompileBlueprintResponse,
    PodAuditSummary,
    PodPropagateResponse,
)


@app.post("/api/pod/1_extract_raw", response_model=PodExtractRawResponse)
async def pod_1_extract_raw(
    file: UploadFile = File(...),
    split_level: int = Form(2),
):
    """
    階段一：初稿生成 API
    上傳 PDF 或 MD 原稿，輸出 Book_raw.docx 與包含幾何座標的 Book_raw_ast.json。
    """
    session_id = f"pod_{uuid.uuid4().hex[:12]}"
    session_dir = PATHS.output_dir / "pod_workspaces" / session_id
    session_dir.mkdir(parents=True, exist_ok=True)

    input_path = session_dir / (file.filename or "input_book.pdf")
    await save_upload_file_safely(file, input_path, max_mb=CFG.max_upload_size_mb)

    ext = input_path.suffix.lower()
    if ext == ".pdf":
        sections, ast_path, raw_docx_path = extract_raw_from_pdf(input_path, session_dir, split_level=split_level)
    elif ext in (".md", ".txt"):
        md_text = input_path.read_text(encoding="utf-8")
        sections, ast_path, raw_docx_path = extract_raw_from_markdown(md_text, session_dir, split_level=split_level)
    else:
        raise HTTPException(status_code=400, detail="生稿僅支援 .pdf 或 .md 檔案")

    return PodExtractRawResponse(
        status="success",
        session_id=session_id,
        raw_docx_name=raw_docx_path.name,
        ast_json_name=ast_path.name,
        total_sections=len(sections),
        sections_preview=[
            PodSectionPreview(id=s.section_id, title=s.title, paragraphs=len(s.paragraphs))
            for s in sections[:10]
        ],
    )


@app.post("/api/pod/2_compile_blueprint", response_model=PodCompileBlueprintResponse)
async def pod_2_compile_blueprint(
    seed_file: UploadFile = File(...),
    pdf_file: Optional[UploadFile] = File(None),
    session_id: Optional[str] = Form(None),
    manual_boundary_anchor: Optional[str] = Form(None),
    manual_boundary_idx: Optional[int] = Form(None),
):
    """
    範本學習與藍圖合成 API
    比對人工精修範本推導樣式定義與狀態機轉換規則。
    若提供原始 PDF (pdf_file)，則直接解析完整幾何結構與三線表，建立 Book_raw_ast.json。
    """
    if not session_id:
        session_id = f"pod_{uuid.uuid4().hex[:12]}"
    session_dir = PATHS.output_dir / "pod_workspaces" / session_id
    session_dir.mkdir(parents=True, exist_ok=True)

    seed_path = session_dir / (seed_file.filename or "Book_seed.docx")
    await save_upload_file_safely(seed_file, seed_path, max_mb=CFG.max_template_size_mb)

    ast_path = session_dir / "Book_raw_ast.json"
    has_source_pdf = False
    detected_tables_count = 0

    if pdf_file and pdf_file.filename:
        source_pdf_path = session_dir / "source_book.pdf"
        await save_upload_file_safely(pdf_file, source_pdf_path, max_mb=CFG.max_upload_size_mb)
        has_source_pdf = True
        try:
            from src.scripts.book_layout_extractor import extract_book_vector_ast
            from src.pod_studio.section_ast import SectionAST
            book_ast = extract_book_vector_ast(source_pdf_path, enable_footnotes=False)
            SectionAST.dump_book_ast_json(book_ast.sections, ast_path)
            detected_tables_count = sum(
                1 for sec in book_ast.sections for p in sec.paragraphs if p.font_name == "table"
            )
        except Exception as e:
            print(f"[POD Warning] Failed to extract vector AST from uploaded source PDF: {e}")

    # When starting directly from docx without prior extraction or PDF, construct paragraph AST
    if not ast_path.exists() and seed_path.suffix.lower() == ".docx":
        import docx
        from src.pod_studio.section_ast import SectionAST, SectionUnit, ParagraphUnit
        seed_doc = docx.Document(str(seed_path))
        sections = [
            SectionUnit(
                section_id=f"sec_{i:03d}",
                title=p.text[:30],
                level=1,
                paragraphs=[ParagraphUnit(para_id=f"p_{i}", raw_text=p.text)],
                raw_content=p.text,
            )
            for i, p in enumerate(seed_doc.paragraphs)
        ]
        SectionAST.dump_book_ast_json(sections, ast_path)

    blueprint = compile_blueprint_from_seed(
        session_dir,
        seed_path,
        manual_boundary_anchor=manual_boundary_anchor,
        manual_boundary_idx=manual_boundary_idx,
    )
    if has_source_pdf:
        blueprint.metadata["has_source_pdf"] = True
        blueprint.metadata["detected_tables_count"] = detected_tables_count
        blueprint.save(session_dir / "book_blueprint.json")

    return PodCompileBlueprintResponse(
        status="success",
        session_id=session_id,
        blueprint=blueprint.to_dict(),
        inferred_rules_count=len(blueprint.style_rules),
        mandatory_badges=blueprint.mandatory_section_badges,
        detected_styles=blueprint.template_styles_used,
        boundary_idx=blueprint.metadata.get("seed_last_modified_idx", 0),
        boundary_source=blueprint.metadata.get("boundary_source", "auto_heuristic"),
        manual_boundary_anchor=blueprint.metadata.get("manual_boundary_anchor", ""),
        boundary_context_text=blueprint.metadata.get("boundary_context_text", ""),
        total_paragraphs=blueprint.metadata.get("seed_total_paragraphs", 0),
        has_source_pdf=has_source_pdf,
        detected_tables_count=detected_tables_count,
    )


@app.post("/api/pod/3_propagate_and_validate", response_model=PodPropagateResponse)
async def pod_3_propagate_and_validate(
    session_id: str = Form(...),
    engine_mode: str = Form("rule_first"),
    template_file: Optional[UploadFile] = File(None),
    blueprint_json: Optional[str] = Form(None),
    model: Optional[str] = Form(None),
    max_continuation_sections: Optional[int] = Form(None),
    manual_boundary_anchor: Optional[str] = Form(None),
    manual_boundary_idx: Optional[int] = Form(None),
):
    """
    階段三：批量擴散與硬性驗收 API (支援 Dual-Engine: rule_first 與 vlm_continuation)
    依藍圖執行確定性 AST 標記索引套用或 VLM 接續，並由 CJK 骨架守恆閘門校驗輸出。
    """
    session_dir = PATHS.output_dir / "pod_workspaces" / session_id
    if not session_dir.exists():
        raise HTTPException(status_code=404, detail="Session 不存在或已過期")

    template_path = None
    if template_file and template_file.filename:
        template_path = session_dir / "template.docx"
        await save_upload_file_safely(template_file, template_path, max_mb=CFG.max_template_size_mb)

    custom_bp = None
    if blueprint_json:
        try:
            custom_bp = BookBlueprint.from_json(blueprint_json)
        except Exception as e:
            raise HTTPException(status_code=400, detail=f"自訂 Blueprint JSON 格式錯誤: {e}")

    pdf_files = list(session_dir.glob("*.pdf"))
    pdf_path = pdf_files[0] if pdf_files else None

    final_docx_path, report, bp = propagate_and_validate_session(
        session_dir=session_dir,
        engine_mode=engine_mode,
        template_docx_path=template_path,
        custom_blueprint=custom_bp,
        pdf_path=pdf_path,
        model=model,
        max_sections=max_continuation_sections,
        manual_boundary_anchor=manual_boundary_anchor,
        manual_boundary_idx=manual_boundary_idx,
    )

    resolved_mode = "vlm_continuation" if engine_mode == "vlm_continuation" else "rule_first"
    return PodPropagateResponse(
        status="success",
        session_id=session_id,
        engine_mode=resolved_mode,
        is_valid=report.is_valid,
        final_docx_name=final_docx_path.name,
        audit_summary=PodAuditSummary(
            total_sections=report.total_sections,
            passed_sections=report.passed_sections,
            failed_sections=report.failed_sections,
            green_count=report.green_count,
            yellow_count=report.yellow_count,
            red_count=report.red_count,
            violations_count=len(report.violations),
        ),
        audit_report=report.to_dict(),
    )


@app.get("/api/pod/download/{session_id}/{filename}")
async def pod_download_file(session_id: str, filename: str):
    """下載 POD 工作區之產出物 (Book_raw.docx, Book_POD_Final.docx, audit_report.json 等)"""
    session_dir = (PATHS.output_dir / "pod_workspaces" / session_id).resolve()
    if not session_dir.exists():
        raise HTTPException(status_code=404, detail="Session 不存在")

    target_file = (session_dir / filename).resolve()
    if not target_file.is_relative_to(session_dir) or not target_file.exists():
        raise HTTPException(status_code=404, detail="檔案不存在或存取受限")

    media_type = "application/octet-stream"
    if filename.endswith(".docx"):
        media_type = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    elif filename.endswith(".json"):
        media_type = "application/json"
    elif filename.endswith(".md"):
        media_type = "text/markdown"

    return FileResponse(path=str(target_file), media_type=media_type, filename=filename)


@app.post("/api/pod/vlm_continuation")
async def pod_vlm_continuation(
    session_id: Optional[str] = Form(None),
    pdf_file: Optional[UploadFile] = File(None),
    seed_file: UploadFile = File(...),
    template_file: Optional[UploadFile] = File(None),
    model_name: str = Form("gemini-2.5-flash"),
    max_continuation_sections: Optional[int] = Form(None),
    split_level: int = Form(2),
    manual_boundary_anchor: Optional[str] = Form(None),
    manual_boundary_idx: Optional[int] = Form(None),
):
    """
    POD 視覺整稿接續 API (VLM Continuation & Infilling)
    結合階段一 PDF 與前 5%~10% 人工精修 DOCX (Golden Seed)，
    利用 Gemini 或 Ollama 視覺多模態模型接續完成全書排版。
    """
    if session_id:
        session_dir = (PATHS.output_dir / "pod_workspaces" / session_id).resolve()
        if not session_dir.exists():
            raise HTTPException(status_code=404, detail="指定 Session 不存在或已過期")
    else:
        session_id = f"pod_vlm_{uuid.uuid4().hex[:12]}"
        session_dir = PATHS.output_dir / "pod_workspaces" / session_id
        session_dir.mkdir(parents=True, exist_ok=True)

    # 1. 處理 PDF 檔案
    pdf_path = None
    if pdf_file and pdf_file.filename:
        pdf_path = session_dir / pdf_file.filename
        await save_upload_file_safely(pdf_file, pdf_path, max_mb=CFG.max_upload_size_mb)
    else:
        existing_pdfs = list(session_dir.glob("*.pdf"))
        if existing_pdfs:
            pdf_path = existing_pdfs[0]

    if not pdf_path or not pdf_path.exists():
        raise HTTPException(status_code=400, detail="請提供原始 PDF 檔案")

    # 2. 確保 AST 存在，若不存在則自 PDF 剖析生稿
    ast_path = session_dir / "Book_raw_ast.json"
    if not ast_path.exists():
        try:
            sections, ast_path, _ = extract_raw_from_pdf(
                pdf_path, session_dir, split_level=split_level
            )
        except Exception as e:
            raise HTTPException(status_code=500, detail=f"PDF 生稿剖析失敗: {e}")

    # 3. 處理前 5%~10% 精修範本檔案 (Seed DOCX)
    seed_path = session_dir / (seed_file.filename or "Book_seed.docx")
    await save_upload_file_safely(seed_file, seed_path, max_mb=CFG.max_template_size_mb)

    # 4. 處理可選的 Template DOCX
    template_path = None
    if template_file and template_file.filename:
        template_path = session_dir / "template.docx"
        await save_upload_file_safely(template_file, template_path, max_mb=CFG.max_template_size_mb)

    # 5. 調度 SecureVLMContinuationEngine
    engine = SecureVLMContinuationEngine(
        workspace_base=PATHS.output_dir / "pod_workspaces"
    )

    try:
        result = await engine.execute_continuation(
            pdf_path=pdf_path,
            raw_ast_path=ast_path,
            seed_docx_path=seed_path,
            template_docx_path=template_path,
            model_name=model_name,
            max_continuation_sections=max_continuation_sections,
            target_session_dir=session_dir,
            manual_boundary_anchor=manual_boundary_anchor,
            manual_boundary_idx=manual_boundary_idx,
        )
    except Exception as e:
        raise HTTPException(status_code=500, detail=f"視覺接續流程執行異常: {str(e)}")

    return {
        "status": "success",
        "session_id": session_id,
        "result": result.to_dict(),
        "download_docx_url": f"/api/pod/download/{session_id}/Book_VLM_Continued.docx",
        "download_md_url": f"/api/pod/download/{session_id}/Book_VLM_Continued.md",
        "audit_report_url": f"/api/pod/download/{session_id}/vlm_audit_report.json",
    }




